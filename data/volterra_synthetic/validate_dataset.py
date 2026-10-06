#!/usr/bin/env python3
"""Independently validate the Volterra synthetic dataset.

Reads ONLY the generated files (CSV, XLSX, PDF), the expected values in
INTERNAL_TEST_ONLY/ground_truth.json and the rule constants in tax_rules.py.
It does not import the generator. Prints a PASS/FAIL report; exit code 1 on failure.

    .venv/bin/python validate_dataset.py
"""

from __future__ import annotations

import csv
import json
import math
import re
import sys
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

import tax_rules as tr

ROOT = Path(__file__).resolve().parent / "volterra_synthetic_data"
TAX_YEAR = tr.TAX_YEAR
RESULTS: list[tuple[str, bool, str]] = []


def check(name: str, ok: bool, detail: str = "") -> None:
    RESULTS.append((name, bool(ok), detail))


def read_csv(rel: str) -> list[dict]:
    with (ROOT / rel).open(newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def read_xlsx(rel: str, sheet: str | None = None) -> list[dict]:
    wb = load_workbook(ROOT / rel, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    return [dict(zip(rows[0], r)) for r in rows[1:]]


def pdf_pages(rel: str) -> list[str]:
    return [" ".join((p.extract_text() or "").split()) for p in PdfReader(str(ROOT / rel)).pages]


def as_date(v) -> date | None:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v)
    if re.fullmatch(r"\d{2}/\d{2}/\d{4}", s):
        return datetime.strptime(s, "%m/%d/%Y").date()
    return date.fromisoformat(s[:10])


def main() -> int:
    gt = json.loads((ROOT / "INTERNAL_TEST_ONLY" / "ground_truth.json").read_text())
    exp = gt["expected"]

    # ---------------------------------------------------------------- files present
    required = ["README.md", "COMPANY/company_and_facility.xlsx", "PLM/product_master.xlsx", "PLM/bom.xlsx",
                "MES/item_master.csv", "MES/production_orders.csv", "MES/material_consumption.csv",
                "ERP/supplier_master.xlsx", "ERP/material_purchase_prices.xlsx", "ERP/sales.csv",
                "INTERNAL_TEST_ONLY/ground_truth.json", "INTERNAL_TEST_ONLY/expected_results.md",
                "INTERNAL_TEST_ONLY/data_dictionary.md", "INTERNAL_TEST_ONLY/intentional_issues.md"]
    missing = [r for r in required if not (ROOT / r).exists()]
    check("All required files present", not missing, ", ".join(missing))

    company = read_xlsx("COMPANY/company_and_facility.xlsx", "LegalEntity")
    facilities = read_xlsx("COMPANY/company_and_facility.xlsx", "Facilities")
    pm = read_xlsx("PLM/product_master.xlsx")
    bom = read_xlsx("PLM/bom.xlsx")
    items = read_csv("MES/item_master.csv")
    orders = read_csv("MES/production_orders.csv")
    cons = read_csv("MES/material_consumption.csv")
    vendors = read_xlsx("ERP/supplier_master.xlsx")
    prices = read_xlsx("ERP/material_purchase_prices.xlsx")
    sales = read_csv("ERP/sales.csv")

    # ---------------------------------------------------------------- company / facility
    ent = {r["FIELD"]: r["VALUE"] for r in company}
    check("Claimant and synthetic EIN present", ent.get("Legal name") == gt["claimant"]["legal_name"] and "SYNTHETIC" in str(ent.get("EIN (SYNTHETIC)")))
    check("Facility in the United States", len(facilities) == 1 and facilities[0]["COUNTRY"] == "United States")
    fac = facilities[0]

    # ---------------------------------------------------------------- identifier mapping
    products = sorted({r["PART_NUMBER"] for r in pm})
    plm_to_erp = {r["PART_NUMBER"]: r["ERP_MATERIAL_NO"] for r in pm}
    mes_to_plm = {r["item_code"]: r["plm_part_ref"] for r in items}
    erp_to_plm = {v: k for k, v in plm_to_erp.items()}
    ok = all(len({p, plm_to_erp[p], next(m for m, q in mes_to_plm.items() if q == p)}) == 3 for p in products)
    check("Product IDs differ across PLM/ERP/MES and reconcile via cross-references (ISS-01)",
          ok and set(mes_to_plm.values()) == set(products) and {o["item_code"] for o in orders} <= set(mes_to_plm), str(products))
    check("Plant codes reconcile (ERP WERKS, MES site -> facility)",
          {o["site"] for o in orders} == {fac["MES_SITE"]} and {s["WERKS"] for s in sales} == {fac["ERP_PLANT"]})
    check("Mappings match ground truth", all(gt["identifier_mappings"][p]["erp_material"] == plm_to_erp[p] for p in products))

    # ---------------------------------------------------------------- BOM references
    price_mat = {r["MATNR"]: r for r in prices}
    vendor_by_id = {r["LIFNR"]: r for r in vendors}
    supplied = {m.split(" ")[0]: v["LIFNR"] for v in vendors for m in v["MATERIALS_SUPPLIED"].split("; ")}
    bad = []
    for r in bom:
        if r["TOP_ASSEMBLY"] not in products or not any(p["PART_NUMBER"] == r["TOP_ASSEMBLY"] and p["REVISION"] == r["TOP_REV"] for p in pm):
            bad.append(f"top {r['TOP_ASSEMBLY']} {r['TOP_REV']}")
        if r["MAKE_BUY"] == "Purchased" and (r["ERP_MATERIAL_NO"] not in price_mat or r["ERP_MATERIAL_NO"] not in supplied):
            bad.append(f"component {r['COMPONENT_ITEM']}")
        if r["BOM_LEVEL"] == 2 and not any(b["COMPONENT_ITEM"] == r["PARENT_ITEM"] and b["TOP_ASSEMBLY"] == r["TOP_ASSEMBLY"]
                                         and b["TOP_REV"] == r["TOP_REV"] and b["BOM_LEVEL"] == 1 for b in bom):
            bad.append(f"parent {r['PARENT_ITEM']}")
    check("BOM references resolve (top assemblies, parents, purchased materials -> price list & vendor)", not bad, "; ".join(bad))
    check("Supplier references resolve (price list vendors exist; every vendor material priced)",
          all(r["LIFNR"] in vendor_by_id for r in prices) and all(m in price_mat for m in supplied), "")

    # ---------------------------------------------------------------- production & consumption
    def rev_on(prod: str, day: date) -> str:
        for r in pm:
            if r["PART_NUMBER"] == prod and as_date(r["EFFECTIVE_FROM"]) <= day and (r["EFFECTIVE_TO"] is None or day <= as_date(r["EFFECTIVE_TO"])):
                return r["REVISION"]
        raise ValueError((prod, day))

    lot_qty = {}
    for c in cons:
        lot_qty.setdefault(c["fg_lot_id"], int(c["fg_lot_qty"]))
    order_info = {}
    for o in orders:
        prod = mes_to_plm[o["item_code"]]
        start = as_date(o["start_ts"])
        order_info[o["lot_id"]] = {"product": prod, "start": start, "end": as_date(o["end_ts"]), "rev": rev_on(prod, start),
                                   "good": int(o["qty_good"]), "scrap": int(o["qty_scrap"]), "planned": int(o["qty_planned"]),
                                   "lot_qty": lot_qty.get(o["lot_id"]), "order": o["order_id"]}
    check("Every production order has material consumption", all(v["lot_qty"] is not None for v in order_info.values()))
    check("Good + scrap <= planned on every order", all(v["good"] + v["scrap"] <= v["planned"] for v in order_info.values()))

    # consumption matches BOM x (lot qty + scrap), with unit conversion
    consumed = defaultdict(float)
    for c in cons:
        consumed[(c["fg_lot_id"], c["component_material"])] += float(c["qty_consumed"])
    mism = []
    for lot, v in order_info.items():
        for r in bom:
            if r["TOP_ASSEMBLY"] == v["product"] and r["TOP_REV"] == v["rev"] and r["MAKE_BUY"] == "Purchased":
                per = r["QTY_PER"] / 1000 if r["UOM"] == "G" else r["QTY_PER"]
                expected_q = per * (v["lot_qty"] + v["scrap"])
                got = consumed.get((lot, r["ERP_MATERIAL_NO"]), 0)
                if not math.isclose(got, expected_q, abs_tol=1e-3):
                    mism.append(f"{lot} {r['ERP_MATERIAL_NO']} {got} vs {expected_q}")
        used = {m for (l, m) in consumed if l == lot}
        allowed = {r["ERP_MATERIAL_NO"] for r in bom if r["TOP_ASSEMBLY"] == v["product"] and r["TOP_REV"] == v["rev"] and r["MAKE_BUY"] == "Purchased"}
        if used - allowed:
            mism.append(f"{lot} consumed off-BOM {used - allowed}")
    check("Material consumption = BOM qty x (finished-lot qty + scrap) for the revision in effect", not mism, "; ".join(mism[:5]))

    for p in products:
        e = exp["production"][p]
        got_rep = sum(v["good"] for v in order_info.values() if v["product"] == p)
        got_true = sum(v["lot_qty"] for v in order_info.values() if v["product"] == p)
        check(f"Production totals {p}", got_rep == e["mes_reported_good"] and got_true == e["finished_lot_good"],
              f"MES {got_rep} / lots {got_true}")

    # ---------------------------------------------------------------- sales
    for s in sales:
        s["_date"] = as_date(s["FKDAT"])
        s["_qty"] = int(s["FKIMG"])
    amount_ok = all(math.isclose(float(s["NETWR"]), s["_qty"] * float(s["NETPR"]), abs_tol=0.005) for s in sales)
    check("Invoice amounts = quantity x unit price", amount_ok)
    for p in products:
        rows = [s for s in sales if erp_to_plm.get(s["MATNR"]) == p and s["_date"].year == TAX_YEAR]
        e = exp["sales"][p]
        check(f"2026 sales totals {p}", sum(s["_qty"] for s in rows) == e["qty_2026"]
              and math.isclose(sum(float(s["NETWR"]) for s in rows), e["revenue_2026"], abs_tol=0.01),
              f"{sum(s['_qty'] for s in rows)} cells")
    batch_bad = [s["VBELN"] for s in sales if s["MATNR"] in erp_to_plm and
                 (s["CHARG"] not in order_info or order_info[s["CHARG"]]["product"] != erp_to_plm[s["MATNR"]])]
    check("Every product sales line references a lot of the same product", not batch_bad, ", ".join(batch_bad))
    shipped = defaultdict(int)
    late = []
    for s in sales:
        if s["CHARG"] in order_info:
            shipped[s["CHARG"]] += s["_qty"]
            if order_info[s["CHARG"]]["end"] >= s["_date"]:
                late.append(s["VBELN"])
    over = [l for l, q in shipped.items() if q > order_info[l]["lot_qty"]]
    check("No lot oversold; every shipment after its lot was completed", not over and not late, f"{over} {late}")
    cum_bad = []
    for p in products:
        for s in sorted([s for s in sales if erp_to_plm.get(s["MATNR"]) == p], key=lambda x: x["_date"]):
            sold = sum(x["_qty"] for x in sales if erp_to_plm.get(x["MATNR"]) == p and x["_date"] <= s["_date"])
            made = sum(v["lot_qty"] for v in order_info.values() if v["product"] == p and v["end"] < s["_date"])
            if sold > made:
                cum_bad.append(f"{p} {s['_date']}")
    check("Cumulative sales <= cumulative finished output at every invoice date", not cum_bad, ", ".join(cum_bad))

    # ---------------------------------------------------------------- quality reports
    reports = {}
    for f in sorted((ROOT / "QUALITY").glob("*.pdf")):
        pages = pdf_pages(f"QUALITY/{f.name}")
        m = re.search(r"Product: (\S+) Revision (\S+)", pages[0])
        cap = re.search(r"Declared rated capacity: ([\d.]+) Ah", pages[2])
        reports[(m.group(1), m.group(2))] = {"file": f.name, "draft": "DRAFT" in pages[0], "rated_ah": float(cap.group(1))}

    # ---------------------------------------------------------------- credit arithmetic
    plm_cap = {(r["PART_NUMBER"]): (r["NOMINAL_VOLTAGE_V"], r["RATED_CAPACITY_AH"]) for r in pm}
    for p in products:
        v, ah_plm = plm_cap[p]
        units = sum(s["_qty"] for s in sales if erp_to_plm.get(s["MATNR"]) == p and s["_date"].year == TAX_YEAR)
        caps = {ah_plm} | {r["rated_ah"] for (rp, _), r in reports.items() if rp == p}
        credits = [tr.cell_credit(units, tr.kwh_per_cell(v, ah)) for ah in caps]
        check(f"Credit arithmetic {p}", math.isclose(min(credits), exp["credit"][p]["low"], abs_tol=0.01)
              and math.isclose(max(credits), exp["credit"][p]["high"], abs_tol=0.01), f"{min(credits):,.2f} - {max(credits):,.2f}")
    total_lo = sum(exp["credit"][p]["low"] for p in products)
    total_hi = sum(exp["credit"][p]["high"] for p in products)
    check("Total credit range matches", math.isclose(total_lo, exp["credit_total"]["low"], abs_tol=0.01)
          and math.isclose(total_hi, exp["credit_total"]["high"], abs_tol=0.01), f"{total_lo:,.2f} - {total_hi:,.2f}")
    for p in products:
        for ah_key, facts in exp["product_facts"][p]["by_capacity_value"].items():
            check(f"{p} meets battery-cell definition at {ah_key} Ah", facts["meets_battery_cell_definition"])

    # ---------------------------------------------------------------- evidence index
    cache: dict[str, list[str]] = {}
    miss = []
    for e in gt["evidence_index"]:
        pages = cache.setdefault(e["file"], pdf_pages(e["file"]))
        if " ".join(e["text"].split()) not in pages[e["page"] - 1]:
            miss.append(f"{e['file']} p{e['page']}: {e['text']}")
    check(f"All {len(gt['evidence_index'])} evidence citations found on the cited PDF page", not miss, "; ".join(miss[:3]))

    # ---------------------------------------------------------------- detect discrepancies independently
    detected: dict[str, list[str]] = defaultdict(list)
    detected["ISS-01"] = products  # identifiers differ by design (verified above)

    supplier_docs = {}
    for f in sorted((ROOT / "SUPPLIER_EVIDENCE").glob("*.pdf")):
        pg = pdf_pages(f"SUPPLIER_EVIDENCE/{f.name}")
        name = re.search(r"Supplier legal name: (.+?) Registered address", pg[0])
        if not name:
            continue  # certificates without ownership/coverage sections
        up = re.search(r"Ultimate parent: (.+?) Coverage", pg[0]).group(1)
        cov = re.search(r"from (\d{4}-\d{2}-\d{2}) through (\d{4}-\d{2}-\d{2})", pg[0])
        supplier_docs.setdefault(name.group(1).strip(), []).append({"file": f.name, "ultimate_parent": up,
                                                                     "from": date.fromisoformat(cov.group(1)), "to": date.fromisoformat(cov.group(2))})
    check("Every vendor has at least one ownership/coverage document", set(supplier_docs) == {v["NAME1"] for v in vendors},
          str(set(v["NAME1"] for v in vendors) - set(supplier_docs)))
    for vname, docs in supplier_docs.items():
        if any(d["ultimate_parent"].lower().startswith("not disclosed") for d in docs):
            detected["ISS-02"].append(vname)
    vendor_name = {v["LIFNR"]: v["NAME1"] for v in vendors}
    for c in cons:
        vname = vendor_name[supplied[c["component_material"]]]
        rec = as_date(c["component_lot_received"])
        if not any(d["from"] <= rec <= d["to"] for d in supplier_docs[vname]):
            if vname not in detected["ISS-03"]:
                detected["ISS-03"].append(vname)
    for p in products:
        revs_2026 = [r for r in pm if r["PART_NUMBER"] == p and as_date(r["EFFECTIVE_FROM"]).year <= TAX_YEAR
                     and (r["EFFECTIVE_TO"] is None or as_date(r["EFFECTIVE_TO"]).year >= TAX_YEAR)]
        if len(revs_2026) > 1:
            detected["ISS-04"].append(p)
    for (p, rev) in sorted({(v["product"], v["rev"]) for v in order_info.values()}):
        if (p, rev) not in reports:
            detected["ISS-05"].append(f"{p} rev {rev}")
    for (p, rev), r in reports.items():
        if not math.isclose(plm_cap[p][1], r["rated_ah"]) or r["draft"]:
            detected["ISS-06"].append(p)
    for lot, v in order_info.items():
        if v["good"] != v["lot_qty"]:
            detected["ISS-07"].append(f"{v['order']} ({v['good'] - v['lot_qty']})")
    for s in sales:
        if s["MATNR"] not in erp_to_plm:
            detected["ISS-08"].append(s["VBELN"])

    documented = {i["id"] for i in gt["intentional_issues"]}
    check("Exactly the documented intentional issues are present (no more, no fewer)",
          set(detected) == documented and len(documented) == 8, f"detected {sorted(detected)} vs documented {sorted(documented)}")
    expect_detail = {"ISS-02": ["Eastbay Separator Co., Ltd."], "ISS-03": ["Nordvik Electrolyte AB"], "ISS-04": ["VX-2170"],
                     "ISS-05": ["VX-2170 rev C"], "ISS-06": ["VX-LFP50"], "ISS-07": ["MO-26-0171 (62)"]}
    for k, want in expect_detail.items():
        check(f"{k} affects exactly {want}", sorted(detected[k]) == sorted(want), str(detected[k]))
    check("ISS-08 is exactly one unmatched invoice line", len(detected["ISS-08"]) == 1, str(detected["ISS-08"]))

    # Lots flagged per issue match the expected lists
    lot_flags = defaultdict(set)
    ely_vendor_docs = supplier_docs["Nordvik Electrolyte AB"]
    for c in cons:
        lot = c["fg_lot_id"]
        if c["component_material"] == next(m for m, v in supplied.items() if vendor_name[v] == "Eastbay Separator Co., Ltd."):
            lot_flags["ISS-02"].add(lot)
        if vendor_name[supplied[c["component_material"]]] == "Nordvik Electrolyte AB":
            rec = as_date(c["component_lot_received"])
            if not any(d["from"] <= rec <= d["to"] for d in ely_vendor_docs):
                lot_flags["ISS-03"].add(lot)
    for lot, v in order_info.items():
        if (v["product"], v["rev"]) not in reports:
            lot_flags["ISS-05"].add(lot)
        if v["product"] == "VX-LFP50":
            lot_flags["ISS-06"].add(lot)
    for k in ("ISS-02", "ISS-03", "ISS-05", "ISS-06"):
        check(f"{k} affected lots match expected results", sorted(lot_flags[k]) == sorted(exp["lots_by_issue"][k]),
              f"{sorted(lot_flags[k])}")

    # Credit on lots with no open issues
    clean = sum(tr.cell_credit(s["_qty"], tr.kwh_per_cell(*plm_cap[erp_to_plm[s["MATNR"]]]))
                for s in sales if s["MATNR"] in erp_to_plm and s["_date"].year == TAX_YEAR
                and not any(s["CHARG"] in lots for lots in lot_flags.values())
                and order_info[s["CHARG"]]["good"] == order_info[s["CHARG"]]["lot_qty"])
    check("Credit on lots with no open issues matches", math.isclose(clean, exp["credit_total"]["no_open_issues"], abs_tol=0.01), f"{clean:,.2f}")

    # ---------------------------------------------------------------- no leakage of private truth
    secrets = [s["true_ultimate_parent"] for s in gt["suppliers"] if s["erp_id"] == "0000100268"] + ["never sent", "double", "ISS-0"]
    leaks = []
    for f in ROOT.rglob("*"):
        if not f.is_file() or "INTERNAL_TEST_ONLY" in f.parts:
            continue
        if f.suffix == ".pdf":
            text = " ".join(pdf_pages(str(f.relative_to(ROOT))))
        elif f.suffix == ".xlsx":
            wb = load_workbook(f, read_only=True)
            text = " ".join(str(c) for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for c in row if c is not None)
        else:
            text = f.read_text(encoding="utf-8")
        leaks += [f"{f.name}: {s}" for s in secrets if s.lower() in text.lower()]
    check("Customer-facing files do not reveal private ground truth or issue labels", not leaks, "; ".join(leaks))

    # ---------------------------------------------------------------- report
    width = max(len(n) for n, _, _ in RESULTS)
    fails = [r for r in RESULTS if not r[1]]
    print("Volterra synthetic dataset - validation report\n")
    for name, ok, detail in RESULTS:
        print(f"  {'PASS' if ok else 'FAIL'}  {name.ljust(width)}  {detail if (detail and not ok) else ''}".rstrip())
    print(f"\n{len(RESULTS) - len(fails)}/{len(RESULTS)} checks passed -> {'PASS' if not fails else 'FAIL'}")
    return 1 if fails else 0


if __name__ == "__main__":
    sys.exit(main())
