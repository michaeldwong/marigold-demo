#!/usr/bin/env python3
"""Generate the Volterra Battery Systems synthetic dataset (tax year 2026).

Everything is derived from the ground-truth definitions in this file, so every
cross-file relationship and every number in the expected results is computed by
code. Re-running the script regenerates identical files.

    .venv/bin/python generate_dataset.py
    .venv/bin/python validate_dataset.py

All companies, people, addresses, identifiers and transactions are FICTIONAL.
"""

from __future__ import annotations

import csv
import json
import math
import shutil
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab import rl_config

rl_config.invariant = 1  # deterministic PDF bytes

from reportlab.lib import colors  # noqa: E402
from reportlab.lib.pagesizes import letter  # noqa: E402
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet  # noqa: E402
from reportlab.lib.units import inch  # noqa: E402
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle  # noqa: E402

import tax_rules as tr

HERE = Path(__file__).resolve().parent
OUT = HERE / "volterra_synthetic_data"
TAX_YEAR = 2026
EXPORT_TS = datetime(2027, 2, 12, 8, 30, 0)  # when the customer pulled the exports


def d(s: str) -> date:
    return date.fromisoformat(s)


# ======================================================================================
# 1. GROUND TRUTH: the true underlying world
# ======================================================================================

COMPANY = {
    "legal_name": "Volterra Battery Systems, Inc.",
    "entity_type": "C corporation",
    "state_of_incorporation": "Delaware",
    "ein": "00-0004521",
    "ein_note": "SYNTHETIC - not a real Employer Identification Number",
    "hq_address": "4100 Lakeshore Industrial Drive, Holland, MI 49423, United States",
    "tax_year": TAX_YEAR,
    "tax_year_start": "2026-01-01",
    "tax_year_end": "2026-12-31",
    "related_entities": [],  # true: no parents, subsidiaries or affiliates
    "claimant_for_45x": True,
}

FACILITY = {
    "facility_id": "FAC-001",
    "erp_plant_code": "MI01",
    "mes_site_code": "HOLLAND-01",
    "name": "Volterra Holland Cell Plant",
    "street": "4100 Lakeshore Industrial Drive",
    "city": "Holland",
    "state": "MI",
    "postal_code": "49423",
    "country": "United States",
    "owner": COMPANY["legal_name"],
    "operator": COMPANY["legal_name"],
    "operations_start": "2026-01-12",
    "lines": [
        {"erp_work_center": "WC-L1", "mes_line": "CYL-01", "description": "Line 1 - cylindrical 21700 cell assembly"},
        {"erp_work_center": "WC-L2", "mes_line": "PRS-01", "description": "Line 2 - prismatic cell assembly"},
    ],
}

# True product facts. For VX-LFP50 the PLM attribute (50.0 Ah) is a stale design target;
# the true rated capacity is 48.6 Ah (the value in the draft test report).
PRODUCTS = {
    "VX-2170": {
        "plm_id": "VX-2170",
        "name": "VX-2170 Cylindrical NMC Cell",
        "component_type": "battery cell",
        "format": "21700 cylindrical",
        "chemistry": "NMC811 cathode / graphite anode",
        "nominal_voltage_v": 3.6,
        "plm_rated_capacity_ah": 5.0,
        "true_rated_capacity_ah": 5.0,
        "diameter_mm": 21.0,
        "height_mm": 70.0,
        "mass_g": 68.5,
        "max_continuous_discharge_a": 20.0,
        "charge_voltage_v": 4.2,
        "cutoff_voltage_v": 2.5,
        "erp_material": "MAT-10004721",
        "erp_description": "CELL 21700 NMC 5.0AH VX-2170",
        "mes_item": "VX2170_R7",
        "mes_description": "21700 NMC CELL R7",
        "mes_routing_version": "R7",
        "line": "L1",
        "revisions": [
            {"rev": "B", "effective_from": "2026-01-15", "effective_to": "2026-06-30", "change_notice": "ECN-2025-118",
             "change_summary": "Production release for Line 1."},
            {"rev": "C", "effective_from": "2026-07-01", "effective_to": None, "change_notice": "ECN-2026-031",
             "change_summary": "Separator changed from SEP-016 (16 um) to SEP-021 (12 um ceramic-coated) for cost reduction."},
        ],
    },
    "VX-LFP50": {
        "plm_id": "VX-LFP50",
        "name": "VX-LFP50 Prismatic LFP Cell",
        "component_type": "battery cell",
        "format": "prismatic, aluminum case",
        "chemistry": "LFP cathode / graphite anode",
        "nominal_voltage_v": 3.2,
        "plm_rated_capacity_ah": 50.0,
        "true_rated_capacity_ah": 48.6,
        "length_mm": 148.0,
        "width_mm": 27.0,
        "height_mm": 101.0,
        "mass_g": 905.0,
        "max_continuous_discharge_a": 50.0,
        "charge_voltage_v": 3.65,
        "cutoff_voltage_v": 2.5,
        "erp_material": "MAT-10004755",
        "erp_description": "CELL PRISM LFP 50AH VX-LFP50",
        "mes_item": "LFP50-PR-A",
        "mes_description": "PRISMATIC LFP CELL 50AH",
        "mes_routing_version": "A",
        "line": "L2",
        "revisions": [
            {"rev": "A", "effective_from": "2026-05-15", "effective_to": None, "change_notice": "ECN-2026-012",
             "change_summary": "Production release for Line 2."},
        ],
    },
}


def cell_volume_l(p: dict) -> float:
    if "diameter_mm" in p:
        return math.pi * (p["diameter_mm"] / 2) ** 2 * p["height_mm"] / 1e6
    return p["length_mm"] * p["width_mm"] * p["height_mm"] / 1e6


# Suppliers. true_* fields are the private truth; customer-facing evidence may say less.
SUPPLIERS = [
    {"erp_id": "0000100231", "name": "Great Plains Cathode LLC", "country": "US", "street": "2200 Prairie Gateway Blvd",
     "city": "Lincoln", "region": "NE", "postal": "68521", "materials": ["MAT-30000102"],
     "procurement_parent_note": "Great Plains Materials Group, Inc. (US)", "true_immediate_parent": "Great Plains Materials Group, Inc.",
     "true_ultimate_parent": "Great Plains Materials Group, Inc.", "true_ultimate_parent_country": "United States", "true_ownership_pct": 100},
    {"erp_id": "0000100236", "name": "Prairie Phosphate Materials Inc.", "country": "US", "street": "88 Industrial Park Road",
     "city": "Pittsburg", "region": "KS", "postal": "66762", "materials": ["MAT-30000108"],
     "procurement_parent_note": "Independent (privately held)", "true_immediate_parent": None,
     "true_ultimate_parent": "Prairie Phosphate Materials Inc. (independent)", "true_ultimate_parent_country": "United States", "true_ownership_pct": None},
    {"erp_id": "0000100244", "name": "Kestrel Graphite Corp.", "country": "CA", "street": "1450 Rue Industrielle",
     "city": "Becancour", "region": "QC", "postal": "G9H 2T4", "materials": ["MAT-30000205"],
     "procurement_parent_note": "Kestrel Resources Ltd. (Canada)", "true_immediate_parent": "Kestrel Resources Ltd.",
     "true_ultimate_parent": "Kestrel Resources Ltd.", "true_ultimate_parent_country": "Canada", "true_ownership_pct": 100},
    {"erp_id": "0000100251", "name": "Tamarack Membrane Inc.", "country": "US", "street": "610 Riverside Commerce Dr",
     "city": "Charlotte", "region": "NC", "postal": "28214", "materials": ["MAT-30000411"],
     "procurement_parent_note": "Independent (privately held)", "true_immediate_parent": None,
     "true_ultimate_parent": "Tamarack Membrane Inc. (independent)", "true_ultimate_parent_country": "United States", "true_ownership_pct": None},
    {"erp_id": "0000100268", "name": "Eastbay Separator Co., Ltd.", "country": "MY", "street": "Lot 21, Jalan Perindustrian 4",
     "city": "Senai", "region": "Johor", "postal": "81400", "materials": ["MAT-30000415"],
     "procurement_parent_note": "Eastbay Holdings Pte. Ltd. (Singapore) per supplier onboarding form",
     "true_immediate_parent": "Eastbay Holdings Pte. Ltd. (Singapore)",
     # PRIVATE TRUTH - not disclosed anywhere in the customer-facing files (issue ISS-02).
     "true_ultimate_parent": "Huaying Advanced Materials Group Co., Ltd.", "true_ultimate_parent_country": "China", "true_ownership_pct": 62},
    {"erp_id": "0000100273", "name": "Nordvik Electrolyte AB", "country": "SE", "street": "Industrivagen 14",
     "city": "Sundsvall", "region": "", "postal": "856 33", "materials": ["MAT-30000520"],
     "procurement_parent_note": "Nordvik Kemi Holding AB (Sweden)", "true_immediate_parent": "Nordvik Kemi Holding AB",
     "true_ultimate_parent": "Nordvik Kemi Holding AB", "true_ultimate_parent_country": "Sweden", "true_ownership_pct": 100},
    {"erp_id": "0000100289", "name": "Lakeshore Can & Stamping Inc.", "country": "US", "street": "905 Harbor Industrial Way",
     "city": "Sandusky", "region": "OH", "postal": "44870", "materials": ["MAT-30000610", "MAT-30000655"],
     "procurement_parent_note": "Lakeshore Industries Holdings LLC (US)", "true_immediate_parent": "Lakeshore Industries Holdings LLC",
     "true_ultimate_parent": "Lakeshore Industries Holdings LLC", "true_ultimate_parent_country": "United States", "true_ownership_pct": 100},
]
SUP = {s["erp_id"]: s for s in SUPPLIERS}

# Materials. base = integer base unit used for exact arithmetic:
#   "dg" (decigram) for materials stocked in KG, "cm2" for M2, "EA" for pieces.
MATERIALS = [
    {"plm_part": "RM-CAM-811", "erp_material": "MAT-30000102", "description": "CATHODE ACTIVE MATERIAL NMC811", "erp_uom": "KG",
     "bom_uom": "G", "base": "dg", "supplier": "0000100231", "price": 38.00, "lot_size_base": 2_500_000, "lot_prefix": "GPC",
     "electrode_active_material": True, "manufacturing_site": "Lincoln, NE, United States"},
    {"plm_part": "RM-CAM-LFP", "erp_material": "MAT-30000108", "description": "CATHODE ACTIVE MATERIAL LFP", "erp_uom": "KG",
     "bom_uom": "G", "base": "dg", "supplier": "0000100236", "price": 9.50, "lot_size_base": 10_000_000, "lot_prefix": "PPM",
     "electrode_active_material": True, "manufacturing_site": "Pittsburg, KS, United States"},
    {"plm_part": "RM-AAM-SG", "erp_material": "MAT-30000205", "description": "ANODE ACTIVE MATERIAL SYNTH GRAPHITE", "erp_uom": "KG",
     "bom_uom": "G", "base": "dg", "supplier": "0000100244", "price": 8.20, "lot_size_base": 5_000_000, "lot_prefix": "KGC",
     "electrode_active_material": True, "manufacturing_site": "Becancour, QC, Canada"},
    {"plm_part": "RM-SEP-016", "erp_material": "MAT-30000411", "description": "SEPARATOR PE 16UM", "erp_uom": "M2",
     "bom_uom": "M2", "base": "cm2", "supplier": "0000100251", "price": 1.10, "lot_size_base": 25_000_000, "lot_prefix": "TMB",
     "electrode_active_material": False, "manufacturing_site": "Charlotte, NC, United States"},
    {"plm_part": "RM-SEP-021", "erp_material": "MAT-30000415", "description": "SEPARATOR PE 12UM CERAMIC COATED", "erp_uom": "M2",
     "bom_uom": "M2", "base": "cm2", "supplier": "0000100268", "price": 0.85, "lot_size_base": 12_000_000, "lot_prefix": "ESC",
     "electrode_active_material": False, "manufacturing_site": "Senai, Johor, Malaysia"},
    # Electrolyte is delivered in drums kitted per production order (dedicated lots).
    {"plm_part": "RM-ELY-LP1", "erp_material": "MAT-30000520", "description": "ELECTROLYTE 1M LIPF6 EC/EMC", "erp_uom": "KG",
     "bom_uom": "G", "base": "dg", "supplier": "0000100273", "price": 9.80, "lot_size_base": None, "drum_kg": 25, "lot_prefix": "NVE",
     "electrode_active_material": False, "manufacturing_site": "Sundsvall, Sweden"},
    {"plm_part": "RM-CAN-217", "erp_material": "MAT-30000610", "description": "CAN+CAP ASSY 21700 NI-PLATED", "erp_uom": "EA",
     "bom_uom": "EA", "base": "EA", "supplier": "0000100289", "price": 0.115, "lot_size_base": 12_000, "lot_prefix": "LCS",
     "electrode_active_material": False, "manufacturing_site": "Sandusky, OH, United States"},
    {"plm_part": "RM-CASE-P50", "erp_material": "MAT-30000655", "description": "PRISMATIC AL CASE+LID 148X27X101", "erp_uom": "EA",
     "bom_uom": "EA", "base": "EA", "supplier": "0000100289", "price": 2.85, "lot_size_base": 3_000, "lot_prefix": "LCP",
     "electrode_active_material": False, "manufacturing_site": "Sandusky, OH, United States"},
]
MAT_BY_PLM = {m["plm_part"]: m for m in MATERIALS}
MAT_BY_ERP = {m["erp_material"]: m for m in MATERIALS}

SUBASSEMBLIES = {
    "SA-CE-2170": "CATHODE ELECTRODE ASSY VX-2170",
    "SA-AE-2170": "ANODE ELECTRODE ASSY VX-2170",
    "SA-CE-LFP50": "CATHODE ELECTRODE ASSY VX-LFP50",
    "SA-AE-LFP50": "ANODE ELECTRODE ASSY VX-LFP50",
}


def bom_lines(product: str, rev: str) -> list[dict]:
    """Two-level BOM: cell -> electrode subassemblies -> purchased materials. Quantities per one cell."""
    if product == "VX-2170":
        sep = "RM-SEP-016" if rev == "B" else "RM-SEP-021"
        spec = [
            (1, "10", "SA-CE-2170", 1, "EA", None), (2, "10.10", "RM-CAM-811", 24.5, "G", "SA-CE-2170"),
            (1, "20", "SA-AE-2170", 1, "EA", None), (2, "20.10", "RM-AAM-SG", 14.2, "G", "SA-AE-2170"),
            (1, "30", sep, 0.081, "M2", None), (1, "40", "RM-ELY-LP1", 5.2, "G", None), (1, "50", "RM-CAN-217", 1, "EA", None),
        ]
    else:
        spec = [
            (1, "10", "SA-CE-LFP50", 1, "EA", None), (2, "10.10", "RM-CAM-LFP", 380, "G", "SA-CE-LFP50"),
            (1, "20", "SA-AE-LFP50", 1, "EA", None), (2, "20.10", "RM-AAM-SG", 165, "G", "SA-AE-LFP50"),
            (1, "30", "RM-SEP-016", 1.15, "M2", None), (1, "40", "RM-ELY-LP1", 140, "G", None), (1, "50", "RM-CASE-P50", 1, "EA", None),
        ]
    return [{"level": lv, "find_no": fn, "component": c, "qty": q, "uom": u, "subassembly_parent": sp} for lv, fn, c, q, u, sp in spec]


def per_cell_base(qty: float, bom_uom: str) -> int:
    """Convert a per-cell BOM quantity to integer base units (g -> dg, m2 -> cm2)."""
    if bom_uom == "G":
        return round(qty * 10)
    if bom_uom == "M2":
        return round(qty * 10_000)
    return int(qty)


def base_to_erp(qty_base: int, base: str) -> float:
    return qty_base / 10_000 if base in ("dg", "cm2") else qty_base


# Production orders. reported_good = MES order confirmation; true_good = finished-lot quantity.
# They differ for exactly one order (issue ISS-07).
ORDERS_SPEC = [
    # product, mes_order, erp_order, start, end, planned, reported_good, true_good, scrap
    ("VX-2170", "MO-26-0101", "1000231", "2026-02-09", "2026-02-20", 5000, 4812, 4812, 188),
    ("VX-2170", "MO-26-0114", "1000248", "2026-03-09", "2026-03-20", 6000, 5874, 5874, 126),
    ("VX-2170", "MO-26-0132", "1000266", "2026-04-06", "2026-04-17", 6500, 6391, 6391, 109),
    ("VX-2170", "MO-26-0157", "1000290", "2026-05-26", "2026-06-12", 7000, 6868, 6868, 132),
    ("VX-LFP50", "MO-26-0163", "1000297", "2026-06-08", "2026-06-19", 1200, 1146, 1146, 54),
    ("VX-2170", "MO-26-0171", "1000305", "2026-07-06", "2026-07-17", 6500, 6212, 6150, 288),
    ("VX-LFP50", "MO-26-0172", "1000306", "2026-07-06", "2026-07-17", 1500, 1452, 1452, 48),
    ("VX-LFP50", "MO-26-0188", "1000322", "2026-08-03", "2026-08-14", 1800, 1761, 1761, 39),
    ("VX-2170", "MO-26-0195", "1000330", "2026-08-17", "2026-08-28", 7000, 6905, 6905, 95),
    ("VX-LFP50", "MO-26-0209", "1000344", "2026-09-08", "2026-09-18", 2000, 1958, 1958, 42),
    ("VX-2170", "MO-26-0221", "1000357", "2026-10-05", "2026-10-16", 7000, 6887, 6887, 113),
    ("VX-LFP50", "MO-26-0222", "1000358", "2026-10-05", "2026-10-16", 2000, 1972, 1972, 28),
    ("VX-LFP50", "MO-26-0240", "1000375", "2026-11-02", "2026-11-13", 2200, 2163, 2163, 37),
    ("VX-2170", "MO-26-0258", "1000391", "2026-11-30", "2026-12-11", 7500, 7322, 7322, 178),
    ("VX-LFP50", "MO-26-0259", "1000392", "2026-12-01", "2026-12-11", 2100, 2040, 2040, 60),
]

CUSTOMERS = {
    "0000300112": {"name": "Ridgeline Power Tools, Inc.", "city": "Columbus", "region": "OH", "related_to_volterra": False},
    "0000300147": {"name": "Northgate E-Bike Co.", "city": "Portland", "region": "OR", "related_to_volterra": False},
    "0000300163": {"name": "Granite Grid Storage LLC", "city": "Austin", "region": "TX", "related_to_volterra": False},
    "0000300188": {"name": "Circuit Loop Recycling LLC", "city": "Grand Rapids", "region": "MI", "related_to_volterra": False},
}

PRICES = {("VX-2170", "0000300112"): 3.85, ("VX-2170", "0000300147"): 3.95, ("VX-LFP50", "0000300163"): 58.00}

# Invoices: date, customer, product, quantity.
SALES_SPEC = [
    ("2026-02-26", "0000300147", "VX-2170", 1000), ("2026-03-10", "0000300112", "VX-2170", 2400),
    ("2026-04-08", "0000300112", "VX-2170", 2800), ("2026-04-22", "0000300147", "VX-2170", 2000),
    ("2026-05-12", "0000300112", "VX-2170", 3000), ("2026-06-09", "0000300112", "VX-2170", 3200),
    ("2026-07-07", "0000300163", "VX-LFP50", 1000), ("2026-07-14", "0000300112", "VX-2170", 3400),
    ("2026-07-28", "0000300147", "VX-2170", 2600), ("2026-08-04", "0000300163", "VX-LFP50", 1300),
    ("2026-08-11", "0000300112", "VX-2170", 3300), ("2026-09-01", "0000300163", "VX-LFP50", 1500),
    ("2026-09-15", "0000300112", "VX-2170", 3600), ("2026-09-22", "0000300147", "VX-2170", 2200),
    ("2026-10-06", "0000300163", "VX-LFP50", 1600), ("2026-10-13", "0000300112", "VX-2170", 3500),
    ("2026-10-27", "0000300147", "VX-2170", 3000), ("2026-11-03", "0000300163", "VX-LFP50", 1700),
    ("2026-11-10", "0000300112", "VX-2170", 3700), ("2026-12-01", "0000300163", "VX-LFP50", 1700),
    ("2026-12-08", "0000300112", "VX-2170", 3600), ("2026-12-21", "0000300147", "VX-2170", 3200),
    ("2026-12-29", "0000300163", "VX-LFP50", 1200),
    # Shipped in January 2027 - present in the ERP extract, outside the 2026 tax year.
    ("2027-01-08", "0000300163", "VX-LFP50", 1500), ("2027-01-12", "0000300112", "VX-2170", 3800),
]
# Ambiguous record (issue ISS-08): B-grade cells sold to a recycler under a material number
# that is not in the PLM product master and with no batch reference.
BGRADE = {"date": "2026-11-24", "customer": "0000300188", "erp_material": "MAT-10004759",
          "description": "VXLFP50 CELL B-GRADE", "qty": 140, "price": 9.00,
          "true_origin": "Downgraded VX-LFP50 cells set aside from 2026 scrap counts after failing outgoing capacity QC."}

# Supplier evidence documents (customer-facing PDFs).
ATTESTATIONS = {
    "0000100231": {"file": "GreatPlainsCathode_Supplier_Attestation_2026.pdf", "kind": "SUPPLIER ATTESTATION",
                   "coverage": ("2026-01-01", "2026-12-31"), "signatory": "Rebecca Lindqvist, Chief Financial Officer",
                   "signed": "2025-12-15", "immediate_parent": "Great Plains Materials Group, Inc. (Delaware, USA) - 100%",
                   "ultimate_parent": "Great Plains Materials Group, Inc. (Delaware, USA)"},
    "0000100236": {"file": "PrairiePhosphate_Supplier_Questionnaire_2026.pdf", "kind": "SUPPLIER QUESTIONNAIRE",
                   "coverage": ("2026-01-01", "2026-12-31"), "signatory": "Thomas Ewing, President",
                   "signed": "2026-04-20", "immediate_parent": "None - privately held by founders",
                   "ultimate_parent": "None - Prairie Phosphate Materials Inc. is not owned by any other entity"},
    "0000100244": {"file": "KestrelGraphite_Supplier_Attestation_2026.pdf", "kind": "SUPPLIER ATTESTATION",
                   "coverage": ("2026-01-01", "2026-12-31"), "signatory": "Marie-Claude Tremblay, VP Finance",
                   "signed": "2025-12-18", "immediate_parent": "Kestrel Resources Ltd. (Ontario, Canada) - 100%",
                   "ultimate_parent": "Kestrel Resources Ltd. (Ontario, Canada); widely held, listed on the TSX"},
    "0000100251": {"file": "TamarackMembrane_Supplier_Attestation_2026.pdf", "kind": "SUPPLIER ATTESTATION",
                   "coverage": ("2026-01-01", "2026-12-31"), "signatory": "Gregory Boateng, Chief Executive Officer",
                   "signed": "2025-12-10", "immediate_parent": "None - privately held",
                   "ultimate_parent": "None - Tamarack Membrane Inc. is not owned by any other entity"},
    "0000100268": {"file": "EastbaySeparator_Supplier_Declaration_2026.pdf", "kind": "SUPPLIER DECLARATION",
                   "coverage": ("2026-06-01", "2027-05-31"), "signatory": "Lim Wei Jie, General Manager",
                   "signed": "2026-05-22", "immediate_parent": "Eastbay Holdings Pte. Ltd. (Singapore) - 100%",
                   "ultimate_parent": "Not disclosed - confidential"},
    "0000100273": {"file": "NordvikElectrolyte_Supplier_Attestation_2025-2026.pdf", "kind": "SUPPLIER ATTESTATION",
                   "coverage": ("2025-07-01", "2026-06-30"), "signatory": "Anders Holm, Chief Financial Officer",
                   "signed": "2025-06-20", "immediate_parent": "Nordvik Kemi Holding AB (Sweden) - 100%",
                   "ultimate_parent": "Nordvik Kemi Holding AB (Sweden)",
                   "true_note": "A renewal for 2026-07-01 to 2027-06-30 was signed by Nordvik but never sent to Volterra."},
    "0000100289": {"file": "LakeshoreCan_Supplier_Attestation_2026.pdf", "kind": "SUPPLIER ATTESTATION",
                   "coverage": ("2026-01-01", "2026-12-31"), "signatory": "Dana Kowalczyk, Controller",
                   "signed": "2025-12-12", "immediate_parent": "Lakeshore Industries Holdings LLC (Ohio, USA) - 100%",
                   "ultimate_parent": "Lakeshore Industries Holdings LLC (Ohio, USA)"},
}
EXTRA_SUPPLIER_DOCS = [
    {"supplier": "0000100244", "file": "KestrelGraphite_Certificate_of_Origin_2026.pdf", "kind": "CERTIFICATE OF ORIGIN"},
]

QUALITY_REPORTS = {
    "VX-2170": {"file": "VX2170_RevB_Qualification_Report.pdf", "doc_no": "QR-2170-B-001", "rev": "B", "status": "APPROVED",
                "test_start": "2026-01-19", "test_end": "2026-01-30", "approved": "2026-02-03",
                "engineer": "Priya Natarajan, Cell Test Engineer", "reviewer": "Daniel Okafor, Quality Manager",
                "measured_ah": [5.04, 5.07, 5.05, 5.08, 5.06, 5.03, 5.09, 5.06, 5.05, 5.07]},
    "VX-LFP50": {"file": "VXLFP50_RevA_Qualification_Report_DRAFT.pdf", "doc_no": "QR-LFP50-A-001", "rev": "A", "status": "DRAFT",
                 "test_start": "2026-05-25", "test_end": "2026-06-03", "approved": None,
                 "engineer": "Marco Bellini, Cell Test Engineer", "reviewer": None,
                 "measured_ah": [48.7, 48.5, 48.6, 48.8, 48.4, 48.6]},
}


# ======================================================================================
# 2. SIMULATION: material lots, consumption, finished lots, sales
# ======================================================================================

def bom_rev_on(product: str, day: date) -> str:
    for r in PRODUCTS[product]["revisions"]:
        if d(r["effective_from"]) <= day and (r["effective_to"] is None or day <= d(r["effective_to"])):
            return r["rev"]
    raise ValueError(f"No BOM revision of {product} effective on {day}")


def simulate():
    orders = []
    for prod, mo, eo, s, e, planned, rep_good, true_good, scrap in ORDERS_SPEC:
        start, end = d(s), d(e)
        prefix = "V21" if prod == "VX-2170" else "LP5"
        orders.append({
            "product": prod, "mes_order": mo, "erp_order": eo, "start": start, "end": end, "planned": planned,
            "reported_good": rep_good, "true_good": true_good, "scrap": scrap, "fg_lot": f"{prefix}{start:%y%m%d}",
            "bom_rev": bom_rev_on(prod, start), "line": PRODUCTS[prod]["line"],
        })
    orders.sort(key=lambda o: (o["start"], o["product"]))

    lots: dict[str, list[dict]] = {m["plm_part"]: [] for m in MATERIALS}
    seq: dict[str, int] = {m["plm_part"]: 0 for m in MATERIALS}
    consumption = []

    def new_lot(mat: dict, size_base: int, received: date) -> dict:
        seq[mat["plm_part"]] += 1
        lot = {"material": mat["plm_part"], "lot": f"{mat['lot_prefix']}{received:%y%m}-{seq[mat['plm_part']]:03d}",
               "received": received, "size_base": size_base, "remaining": size_base, "supplier": mat["supplier"]}
        lots[mat["plm_part"]].append(lot)
        return lot

    for o in orders:
        units_in = o["true_good"] + o["scrap"]  # material is consumed by good and scrapped cells alike
        for line in bom_lines(o["product"], o["bom_rev"]):
            if line["component"] not in MAT_BY_PLM:
                continue  # subassemblies are made in-house; consumption is backflushed to purchased materials
            mat = MAT_BY_PLM[line["component"]]
            need = per_cell_base(line["qty"], line["uom"]) * units_in
            receipt = o["start"] - timedelta(days=5)
            if mat["lot_size_base"] is None:  # dedicated drums for this order
                drums = math.ceil(need / (mat["drum_kg"] * 10_000))
                pool = [new_lot(mat, drums * mat["drum_kg"] * 10_000, receipt)]
            else:
                pool = [l for l in lots[mat["plm_part"]] if l["remaining"] > 0]
            while need > 0:
                if not pool:
                    pool = [new_lot(mat, mat["lot_size_base"], receipt)]
                lot = pool[0]
                take = min(need, lot["remaining"])
                lot["remaining"] -= take
                need -= take
                consumption.append({"order": o, "material": mat, "lot": lot, "qty_base": take, "posted": o["end"]})
                if lot["remaining"] == 0:
                    pool.pop(0)
                if mat["lot_size_base"] is None and need == 0:
                    break

    # Sales: allocate each invoice FIFO (first-in, first-out) to finished lots completed before the invoice date.
    # An invoice line is created per lot drawn from, so some invoices have two lines.
    remaining = {o["fg_lot"]: o["true_good"] for o in orders}
    sales = []
    inv = 90041207
    for s_date, cust, prod, qty in sorted(SALES_SPEC):
        inv += 1
        day = d(s_date)
        need = qty
        posnr = 0
        for o in [o for o in orders if o["product"] == prod and o["end"] < day]:
            if need == 0:
                break
            take = min(need, remaining[o["fg_lot"]])
            if take <= 0:
                continue
            remaining[o["fg_lot"]] -= take
            need -= take
            posnr += 10
            sales.append({"invoice": str(inv), "item": posnr, "date": day, "customer": cust, "product": prod,
                          "erp_material": PRODUCTS[prod]["erp_material"], "description": PRODUCTS[prod]["erp_description"],
                          "batch": o["fg_lot"], "qty": take, "price": PRICES[(prod, cust)]})
        if need:
            raise AssertionError(f"Invoice {inv} oversells {prod} by {need}")
    inv += 1
    sales.append({"invoice": str(inv), "item": 10, "date": d(BGRADE["date"]), "customer": BGRADE["customer"], "product": None,
                  "erp_material": BGRADE["erp_material"], "description": BGRADE["description"], "batch": "",
                  "qty": BGRADE["qty"], "price": BGRADE["price"]})
    sales.sort(key=lambda r: (r["date"], r["invoice"], r["item"]))
    return orders, lots, consumption, sales, remaining


# ======================================================================================
# 3. EXPECTED RESULTS (computed)
# ======================================================================================

ISSUES = {
    "ISS-01": "Product and site identifiers differ across PLM, MES and ERP",
    "ISS-02": "Eastbay Separator ultimate parent not disclosed",
    "ISS-03": "Nordvik electrolyte attestation expired 2026-06-30, no renewal on file",
    "ISS-04": "VX-2170 BOM revision C effective 2026-07-01 (separator change)",
    "ISS-05": "No qualification report covers VX-2170 revision C",
    "ISS-06": "VX-LFP50 rated capacity conflict: PLM 50.0 Ah vs draft test report 48.6 Ah",
    "ISS-07": "MES order confirmation exceeds finished-lot quantity for MO-26-0171 by 62 cells",
    "ISS-08": "B-grade invoice line cannot be linked to an eligible product or production lot",
}


def compute_expected(orders, lots, consumption, sales):
    nordvik_end = d(ATTESTATIONS["0000100273"]["coverage"][1])
    exp: dict = {}

    # Product facts
    prod_facts = {}
    for pid, p in PRODUCTS.items():
        vol = cell_volume_l(p)
        ah_values = sorted({p["plm_rated_capacity_ah"], p["true_rated_capacity_ah"]})
        per_value = {}
        for ah in ah_values:
            wh = p["nominal_voltage_v"] * ah
            power_w = p["nominal_voltage_v"] * p["max_continuous_discharge_a"]
            checks = tr.cell_definition_checks(wh, wh / vol, wh / power_w)
            per_value[str(ah)] = {"energy_wh": round(wh, 3), "energy_density_wh_per_l": round(wh / vol, 1),
                                  "capacity_to_power_ratio": round(wh / power_w, 3),
                                  "kwh_per_cell": round(tr.kwh_per_cell(p["nominal_voltage_v"], ah), 5),
                                  "definition_checks": checks, "meets_battery_cell_definition": all(checks.values())}
        prod_facts[pid] = {"volume_l": round(vol, 5), "by_capacity_value": per_value}
    exp["product_facts"] = prod_facts

    # Lot-level open issues (rules applied to customer-facing facts only)
    lot_issues: dict[str, list[str]] = {}
    for o in orders:
        iss = []
        used = [c for c in consumption if c["order"] is o]
        if any(c["material"]["plm_part"] == "RM-SEP-021" for c in used):
            iss.append("ISS-02")
        if any(c["material"]["plm_part"] == "RM-ELY-LP1" and c["lot"]["received"] > nordvik_end for c in used):
            iss.append("ISS-03")
        if o["product"] == "VX-2170" and o["bom_rev"] == "C":
            iss.append("ISS-05")
        if o["product"] == "VX-LFP50":
            iss.append("ISS-06")
        if o["reported_good"] != o["true_good"]:
            iss.append("ISS-07")
        lot_issues[o["fg_lot"]] = iss
    exp["lot_open_issues"] = lot_issues

    # Production
    prod = {}
    for pid in PRODUCTS:
        os_ = [o for o in orders if o["product"] == pid]
        prod[pid] = {"orders": len(os_), "mes_reported_good": sum(o["reported_good"] for o in os_),
                     "finished_lot_good": sum(o["true_good"] for o in os_), "scrap": sum(o["scrap"] for o in os_),
                     "planned": sum(o["planned"] for o in os_)}
    exp["production"] = prod

    # Sales
    in_year = [s for s in sales if s["date"].year == TAX_YEAR]
    sales_sum = {}
    for pid in PRODUCTS:
        rows = [s for s in in_year if s["product"] == pid]
        nxt = [s for s in sales if s["product"] == pid and s["date"].year > TAX_YEAR]
        sales_sum[pid] = {"invoices_2026": len({s["invoice"] for s in rows}), "lines_2026": len(rows),
                          "qty_2026": sum(s["qty"] for s in rows), "revenue_2026": round(sum(s["qty"] * s["price"] for s in rows), 2),
                          "qty_excluded_2027": sum(s["qty"] for s in nxt), "invoices_excluded_2027": sorted({s["invoice"] for s in nxt}),
                          "ending_inventory_2026_12_31": prod[pid]["finished_lot_good"] - sum(s["qty"] for s in rows)}
    bg = [s for s in sales if s["product"] is None]
    exp["sales"] = sales_sum
    exp["excluded_unmatched_lines"] = [{"invoice": s["invoice"], "erp_material": s["erp_material"], "qty": s["qty"],
                                        "amount": round(s["qty"] * s["price"], 2)} for s in bg]
    exp["sales_counts"] = {"invoices_total": len({s["invoice"] for s in sales}), "lines_total": len(sales),
                           "invoices_2026": len({s["invoice"] for s in in_year}), "lines_2026": len(in_year)}

    # Credit
    credit = {}
    for pid, p in PRODUCTS.items():
        rows = [s for s in in_year if s["product"] == pid]
        ah_vals = {"plm_value": p["plm_rated_capacity_ah"], "test_report_value": p["true_rated_capacity_ah"]}
        by_basis = {}
        for basis, ah in ah_vals.items():
            kwh_u = tr.kwh_per_cell(p["nominal_voltage_v"], ah)
            total = tr.cell_credit(sum(s["qty"] for s in rows), kwh_u)
            clean_qty = sum(s["qty"] for s in rows if not lot_issues[s["batch"]])
            by_basis[basis] = {"rated_capacity_ah": ah, "kwh_per_cell": round(kwh_u, 5),
                               "kwh_total": round(sum(s["qty"] for s in rows) * kwh_u, 3), "credit": total,
                               "credit_no_open_issues": tr.cell_credit(clean_qty, kwh_u),
                               "credit_with_open_issues": round(total - tr.cell_credit(clean_qty, kwh_u), 2)}
        conflict = p["plm_rated_capacity_ah"] != p["true_rated_capacity_ah"]
        credit[pid] = {"units_sold_2026": sum(s["qty"] for s in rows), "capacity_conflict": conflict,
                       "by_capacity_basis": by_basis if conflict else {"plm_value": by_basis["plm_value"]},
                       "low": min(v["credit"] for v in by_basis.values()), "high": max(v["credit"] for v in by_basis.values())}
    exp["credit"] = credit
    exp["credit_total"] = {"low": round(sum(c["low"] for c in credit.values()), 2),
                           "high": round(sum(c["high"] for c in credit.values()), 2),
                           "no_open_issues": round(sum(c["by_capacity_basis"]["plm_value"]["credit_no_open_issues"] for c in credit.values()), 2)}

    # Credit attributed by issue (informational): each sale line's credit listed under each open issue on its lot.
    by_issue = {}
    for s in in_year:
        if s["product"] is None:
            continue
        p = PRODUCTS[s["product"]]
        c = tr.cell_credit(s["qty"], tr.kwh_per_cell(p["nominal_voltage_v"], p["plm_rated_capacity_ah"]))
        for i in lot_issues[s["batch"]]:
            by_issue[i] = round(by_issue.get(i, 0) + c, 2)
    exp["credit_touched_by_issue_plm_basis"] = by_issue

    # Material-assistance inputs (NOT a pass/fail)
    macr = {}
    for pid, p in PRODUCTS.items():
        for r in p["revisions"]:
            cost = 0.0
            lines = []
            for ln in bom_lines(pid, r["rev"]):
                if ln["component"] not in MAT_BY_PLM:
                    continue
                m = MAT_BY_PLM[ln["component"]]
                erp_qty = base_to_erp(per_cell_base(ln["qty"], ln["uom"]), m["base"])
                c = round(erp_qty * m["price"], 6)
                cost += c
                lines.append({"material": m["erp_material"], "supplier": SUP[m["supplier"]]["name"], "qty_per_cell": erp_qty,
                              "uom": m["erp_uom"], "unit_price": m["price"], "cost_per_cell": round(c, 4)})
            macr[f"{pid} rev {r['rev']}"] = {"direct_material_cost_per_cell": round(cost, 4), "lines": lines}
    spend = {}
    for c in consumption:
        sid = c["material"]["supplier"]
        spend[sid] = spend.get(sid, 0) + base_to_erp(c["qty_base"], c["material"]["base"]) * c["material"]["price"]
    exp["material_assistance_inputs"] = {"status": "NOT COMPUTED - requires professional review (tax_rules.UNRESOLVED_RULES U1)",
                                         "cost_per_cell_by_bom_revision": macr,
                                         "consumed_material_cost_2026_by_supplier": {SUP[k]["name"]: round(v, 2) for k, v in sorted(spend.items())}}

    # Affected lots per issue
    exp["lots_by_issue"] = {i: [lot for lot, iss in lot_issues.items() if i in iss] for i in ISSUES}
    return exp


# ======================================================================================
# 4. FILE WRITERS
# ======================================================================================

HDR_FILL = PatternFill("solid", fgColor="DDE3EA")


def write_xlsx(path: Path, sheets: list[tuple[str, list[str], list[list]]]) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    for title, header, rows in sheets:
        ws = wb.create_sheet(title)
        ws.append(header)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = HDR_FILL
            c.alignment = Alignment(vertical="top", wrap_text=True)
        for r in rows:
            ws.append(r)
        for i, h in enumerate(header, start=1):
            width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows if r[i - 1] is not None])
            ws.column_dimensions[get_column_letter(i)].width = min(60, width + 2)
        ws.freeze_panes = "A2"
    wb.properties.creator = "export"
    wb.properties.created = EXPORT_TS
    wb.properties.modified = EXPORT_TS
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_csv(path: Path, header: list[str], rows: list[list]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)


def fmt_qty(q: float, uom: str) -> str:
    return f"{q:.4f}" if uom in ("KG", "M2") else str(int(q))


def write_plm(orders):
    rows = []
    for pid, p in PRODUCTS.items():
        for r in p["revisions"]:
            dims = (f"D{p['diameter_mm']:.1f} x H{p['height_mm']:.1f}" if "diameter_mm" in p
                    else f"L{p['length_mm']:.1f} x W{p['width_mm']:.1f} x H{p['height_mm']:.1f}")
            state = "Released" if r["effective_to"] is None else "Superseded"
            rows.append([pid, r["rev"], p["name"], f"{p['format']} cell, {p['chemistry']}", state,
                         d(r["effective_from"]), d(r["effective_to"]) if r["effective_to"] else None, r["change_notice"],
                         r["change_summary"], p["erp_material"], "Battery cell", p["format"], p["chemistry"],
                         p["nominal_voltage_v"], p["plm_rated_capacity_ah"], dims, p["mass_g"], p["max_continuous_discharge_a"],
                         p["charge_voltage_v"], p["cutoff_voltage_v"], f"ES-{pid.split('-')[1]}-{r['rev']}"])
    header = ["PART_NUMBER", "REVISION", "PART_NAME", "DESCRIPTION", "LIFECYCLE_STATE", "EFFECTIVE_FROM", "EFFECTIVE_TO",
              "CHANGE_NOTICE", "CHANGE_DESCRIPTION", "ERP_MATERIAL_NO", "ITEM_CLASS", "FORM_FACTOR", "CHEMISTRY",
              "NOMINAL_VOLTAGE_V", "RATED_CAPACITY_AH", "DIMENSIONS_MM", "MASS_G", "MAX_CONT_DISCHARGE_A",
              "CHARGE_VOLTAGE_V", "DISCHARGE_CUTOFF_V", "SPEC_DOC"]
    write_xlsx(OUT / "PLM" / "product_master.xlsx", [("Products", header, rows)])

    brow = []
    for pid, p in PRODUCTS.items():
        for r in p["revisions"]:
            for ln in bom_lines(pid, r["rev"]):
                comp = ln["component"]
                m = MAT_BY_PLM.get(comp)
                desc = m["description"] if m else SUBASSEMBLIES[comp]
                parent = ln["subassembly_parent"] or pid
                brow.append([pid, r["rev"], ln["level"], ln["find_no"], parent, comp, desc, ln["qty"], ln["uom"],
                             "Purchased" if m else "Make", m["erp_material"] if m else None,
                             d(r["effective_from"]), d(r["effective_to"]) if r["effective_to"] else None, r["change_notice"]])
    header = ["TOP_ASSEMBLY", "TOP_REV", "BOM_LEVEL", "FIND_NO", "PARENT_ITEM", "COMPONENT_ITEM", "COMPONENT_DESCRIPTION",
              "QTY_PER", "UOM", "MAKE_BUY", "ERP_MATERIAL_NO", "EFFECTIVE_FROM", "EFFECTIVE_TO", "CHANGE_NOTICE"]
    write_xlsx(OUT / "PLM" / "bom.xlsx", [("BOM", header, brow)])


def write_mes(orders, consumption):
    items = [[p["mes_item"], p["mes_description"], pid, p["mes_routing_version"],
              next(l["mes_line"] for l in FACILITY["lines"] if l["erp_work_center"] == "WC-" + p["line"]), "EA", "ACTIVE"]
             for pid, p in PRODUCTS.items()]
    write_csv(OUT / "MES" / "item_master.csv", ["item_code", "item_description", "plm_part_ref", "routing_version", "default_line", "uom", "status"], items)

    line_code = {l["erp_work_center"]: l["mes_line"] for l in FACILITY["lines"]}
    rows = []
    for o in orders:
        p = PRODUCTS[o["product"]]
        rows.append([o["mes_order"], o["erp_order"], p["mes_item"], FACILITY["mes_site_code"], line_code["WC-" + o["line"]],
                     o["fg_lot"], f"{o['start']}T06:00:00", f"{o['end']}T22:00:00", o["planned"], o["reported_good"], o["scrap"],
                     "EA", "CLOSED"])
    write_csv(OUT / "MES" / "production_orders.csv",
              ["order_id", "erp_order_ref", "item_code", "site", "line", "lot_id", "start_ts", "end_ts", "qty_planned",
               "qty_good", "qty_scrap", "uom", "order_status"], rows)

    crow = []
    for c in sorted(consumption, key=lambda c: (c["order"]["start"], c["order"]["mes_order"], c["material"]["erp_material"], c["lot"]["lot"])):
        o, m = c["order"], c["material"]
        crow.append([o["mes_order"], o["fg_lot"], o["true_good"], m["erp_material"], c["lot"]["lot"], c["lot"]["received"].isoformat(),
                     fmt_qty(base_to_erp(c["qty_base"], m["base"]), m["erp_uom"]), m["erp_uom"], f"{c['posted']}T22:15:00"])
    write_csv(OUT / "MES" / "material_consumption.csv",
              ["order_id", "fg_lot_id", "fg_lot_qty", "component_material", "component_lot", "component_lot_received",
               "qty_consumed", "uom", "posted_ts"], crow)


def write_erp(sales):
    srows = []
    for s in SUPPLIERS:
        mats = "; ".join(f"{m} {MAT_BY_ERP[m]['description']}" for m in s["materials"])
        srows.append([s["erp_id"], s["name"], s["country"], s["street"], s["city"], s["region"], s["postal"], mats,
                      s["procurement_parent_note"], "NET60" if s["country"] == "US" else "NET90", "Active"])
    write_xlsx(OUT / "ERP" / "supplier_master.xlsx", [("Vendors", ["LIFNR", "NAME1", "LAND1", "STRAS", "ORT01", "REGIO", "PSTLZ",
                                                                  "MATERIALS_SUPPLIED", "PARENT_COMPANY_NOTE", "ZTERM", "STATUS"], srows)])
    prow = [[m["erp_material"], m["description"], m["erp_uom"], m["price"], "USD", m["supplier"], SUP[m["supplier"]]["name"],
             date(2026, 1, 1), date(2026, 12, 31)] for m in MATERIALS]
    write_xlsx(OUT / "ERP" / "material_purchase_prices.xlsx",
               [("PriceList", ["MATNR", "MAKTX", "MEINS", "NETPR", "WAERS", "LIFNR", "VENDOR_NAME", "VALID_FROM", "VALID_TO"], prow)])

    hdr = ["VBELN", "POSNR", "FKDAT", "KUNNR", "NAME1", "MATNR", "ARKTX", "CHARG", "FKIMG", "VRKME", "NETPR", "NETWR", "WAERK", "WERKS", "LAND1"]
    rows = []
    for s in sales:
        rows.append([s["invoice"], f"{s['item']:06d}", s["date"].strftime("%m/%d/%Y"), s["customer"], CUSTOMERS[s["customer"]]["name"],
                     s["erp_material"], s["description"], s["batch"], s["qty"], "EA", f"{s['price']:.2f}", f"{s['qty'] * s['price']:.2f}",
                     "USD", FACILITY["erp_plant_code"], "US"])
    write_csv(OUT / "ERP" / "sales.csv", hdr, rows)


def write_company():
    c = COMPANY
    write_xlsx(OUT / "COMPANY" / "company_and_facility.xlsx", [
        ("LegalEntity", ["FIELD", "VALUE"], [
            ["Legal name", c["legal_name"]], ["Entity type", c["entity_type"]], ["State of incorporation", c["state_of_incorporation"]],
            ["EIN (SYNTHETIC)", f"{c['ein']} - SYNTHETIC, NOT A REAL EIN"], ["Headquarters address", c["hq_address"]],
            ["Tax year", str(c["tax_year"])], ["Tax year start", c["tax_year_start"]], ["Tax year end", c["tax_year_end"]],
            ["Parent company", "None"], ["Subsidiaries / affiliates", "None"]]),
        ("Facilities", ["FACILITY_ID", "FACILITY_NAME", "ERP_PLANT", "MES_SITE", "STREET", "CITY", "STATE", "ZIP", "COUNTRY",
                        "OWNER", "OPERATOR", "OPERATIONS_START"],
         [[FACILITY["facility_id"], FACILITY["name"], FACILITY["erp_plant_code"], FACILITY["mes_site_code"], FACILITY["street"],
           FACILITY["city"], FACILITY["state"], FACILITY["postal_code"], FACILITY["country"], FACILITY["owner"], FACILITY["operator"],
           FACILITY["operations_start"]]]),
        ("ProductionLines", ["FACILITY_ID", "ERP_WORK_CENTER", "MES_LINE", "DESCRIPTION"],
         [[FACILITY["facility_id"], l["erp_work_center"], l["mes_line"], l["description"]] for l in FACILITY["lines"]]),
    ])


# ---------------------------------- PDFs ---------------------------------------------

STY = getSampleStyleSheet()
H1 = ParagraphStyle("h1", parent=STY["Heading1"], fontSize=15, spaceAfter=6)
H2 = ParagraphStyle("h2", parent=STY["Heading2"], fontSize=11.5, spaceBefore=8, spaceAfter=4)
BODY = ParagraphStyle("body", parent=STY["BodyText"], fontSize=9.5, leading=13)
SMALL = ParagraphStyle("small", parent=BODY, fontSize=8, textColor=colors.HexColor("#555555"))


def kv(label: str, value: str) -> Paragraph:
    return Paragraph(f"<b>{label}:</b> {value}", BODY)


def simple_table(rows: list[list[str]], widths=None) -> Table:
    t = Table(rows, colWidths=widths, hAlign="LEFT")
    t.setStyle(TableStyle([("FONT", (0, 0), (-1, -1), "Helvetica", 8.5), ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 8.5),
                           ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E6EAF0")), ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                           ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    return t


def build_pdf(path: Path, header: str, doc_no: str, pages: list[list], watermark: str | None = None) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)

    def deco(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.drawString(0.75 * inch, 10.6 * inch, header)
        canvas.drawRightString(7.75 * inch, 10.6 * inch, doc_no)
        canvas.drawString(0.75 * inch, 0.5 * inch, "FICTIONAL DOCUMENT - synthetic test data")
        canvas.drawRightString(7.75 * inch, 0.5 * inch, f"Page {doc.page}")
        if watermark:
            canvas.setFont("Helvetica-Bold", 54)
            canvas.setFillColor(colors.Color(0.85, 0.2, 0.2, alpha=0.18))
            canvas.translate(4.25 * inch, 5.5 * inch)
            canvas.rotate(35)
            canvas.drawCentredString(0, 0, watermark)
        canvas.restoreState()

    story = []
    for i, page in enumerate(pages):
        if i:
            story.append(PageBreak())
        story.extend(page)
    doc = SimpleDocTemplate(str(path), pagesize=letter, topMargin=0.9 * inch, bottomMargin=0.8 * inch,
                            leftMargin=0.75 * inch, rightMargin=0.75 * inch, title=doc_no, author="synthetic", creator="synthetic")
    doc.build(story, onFirstPage=deco, onLaterPages=deco)


EVIDENCE_INDEX: list[dict] = []


def ev(fact: str, file: str, page: int, text: str) -> str:
    EVIDENCE_INDEX.append({"fact": fact, "file": file, "page": page, "text": text})
    return text


def write_quality():
    for pid, q in QUALITY_REPORTS.items():
        p = PRODUCTS[pid]
        rel = f"QUALITY/{q['file']}"
        ah = p["true_rated_capacity_ah"]
        vol = cell_volume_l(p)
        wh = p["nominal_voltage_v"] * ah
        mean = sum(q["measured_ah"]) / len(q["measured_ah"])
        power = p["nominal_voltage_v"] * p["max_continuous_discharge_a"]
        dims = (f"Diameter {p['diameter_mm']:.1f} mm, height {p['height_mm']:.1f} mm" if "diameter_mm" in p
                else f"{p['length_mm']:.1f} mm x {p['width_mm']:.1f} mm x {p['height_mm']:.1f} mm")
        status = "Approved" if q["status"] == "APPROVED" else "DRAFT - not approved"
        p1 = [Paragraph("Volterra Battery Systems, Inc. - Cell Qualification Test Report", H1),
              kv("Document number", q["doc_no"]), kv("Document status", status),
              Paragraph(ev(f"{pid} report scope", rel, 1, f"Product: {pid} Revision {q['rev']}"), BODY),
              kv("Product name", p["name"]),
              kv("Test period", f"{q['test_start']} to {q['test_end']}"),
              kv("Test site", "Volterra Holland Cell Plant - Cell Test Laboratory"),
              kv("Test procedures", "IEC 61960-3 (rated capacity), Volterra TP-CELL-004 (dimensions), TP-CELL-007 (rate capability)"),
              Spacer(1, 8), Paragraph("Summary", H2),
              simple_table([["Item", "Result"], ["Rated capacity (declared)", f"{ah:.1f} Ah"],
                            ["Nominal voltage", f"{p['nominal_voltage_v']:.1f} V"],
                            ["Rated energy", f"{wh:.2f} Wh"], ["Result against internal specification", "Meets specification" if q["status"] == "APPROVED" else "Pending approval"]],
                           [3.2 * inch, 2.6 * inch]),
              Spacer(1, 8), Paragraph(f"This report applies only to {pid} Revision {q['rev']} as built on the qualification lot.", SMALL)]
        p2 = [Paragraph("1. Cell identification and dimensions", H2),
              kv("Form factor", p["format"]), kv("Chemistry", p["chemistry"]),
              Paragraph(ev(f"{pid} nominal voltage", rel, 2, f"Nominal voltage: {p['nominal_voltage_v']:.1f} V"), BODY),
              kv("Charge voltage", f"{p['charge_voltage_v']} V"), kv("Discharge cut-off voltage", f"{p['cutoff_voltage_v']} V"),
              kv("Measured external dimensions", dims),
              Paragraph(ev(f"{pid} external volume", rel, 2, f"External volume: {vol * 1000:.2f} mL"), BODY),
              kv("Mass (mean of samples)", f"{p['mass_g']:.1f} g"),
              Spacer(1, 6), Paragraph("External volume is calculated from measured external dimensions excluding terminals.", SMALL)]
        sample_rows = [["Sample", "Discharge capacity (Ah)"]] + [[f"S{i + 1:02d}", f"{v:.2f}"] for i, v in enumerate(q["measured_ah"])]
        p3 = [Paragraph("2. Capacity and energy", H2),
              kv("Procedure", "Charge CC-CV to charge voltage at 25 C; rest 1 h; discharge at 0.2C to cut-off voltage"),
              simple_table(sample_rows, [1.2 * inch, 2.2 * inch]), Spacer(1, 6),
              kv("Mean measured discharge capacity", f"{mean:.2f} Ah"),
              Paragraph(ev(f"{pid} rated capacity", rel, 3, f"Declared rated capacity: {ah:.1f} Ah"), BODY),
              Paragraph(ev(f"{pid} rated energy", rel, 3, f"Rated energy (nominal voltage x rated capacity): {wh:.2f} Wh"), BODY),
              Paragraph(ev(f"{pid} energy density", rel, 3, f"Volumetric energy density (rated energy / external volume): {wh / vol:.0f} Wh/L"), BODY)]
        signoff = ([kv("Test engineer", f"{q['engineer']} (signed {q['test_end']})"),
                    kv("Reviewer", f"{q['reviewer']} (signed {q['approved']})"), kv("Approval date", q["approved"])]
                   if q["status"] == "APPROVED" else
                   [kv("Test engineer", f"{q['engineer']} (signed {q['test_end']})"),
                    kv("Reviewer", "________________ (pending)"), kv("Approval date", "________________")])
        p4 = [Paragraph("3. Rate capability", H2),
              kv("Maximum continuous discharge current", f"{p['max_continuous_discharge_a']:.0f} A"),
              kv("Maximum continuous discharge power at nominal voltage", f"{power:.0f} W"),
              Paragraph(ev(f"{pid} capacity-to-power ratio", rel, 4, f"Capacity-to-power ratio (rated energy / max continuous power): {wh / power:.2f} : 1"), BODY),
              Spacer(1, 10), Paragraph("4. Sign-off", H2), *signoff]
        build_pdf(OUT / rel, "Volterra Battery Systems, Inc. - Quality", q["doc_no"], [p1, p2, p3, p4],
                  watermark="DRAFT" if q["status"] == "DRAFT" else None)


def write_supplier_docs():
    for sid, a in ATTESTATIONS.items():
        s = SUP[sid]
        rel = f"SUPPLIER_EVIDENCE/{a['file']}"
        mats = [MAT_BY_ERP[m] for m in s["materials"]]
        addr = f"{s['street']}, {s['city']}{', ' + s['region'] if s['region'] else ''} {s['postal']}, {s['country']}"
        p1 = [Paragraph(f"{a['kind'].title()} to Volterra Battery Systems, Inc.", H1),
              Paragraph(ev(f"{s['name']} legal name", rel, 1, f"Supplier legal name: {s['name']}"), BODY),
              kv("Registered address", addr),
              kv("Volterra vendor number", sid),
              Spacer(1, 6), Paragraph("Materials covered", H2),
              simple_table([["Volterra material", "Description", "Manufacturing site"]] +
                           [[m["erp_material"], m["description"], m["manufacturing_site"]] for m in mats],
                           [1.3 * inch, 2.9 * inch, 2.4 * inch]),
              Spacer(1, 6), Paragraph("Ownership", H2),
              Paragraph(ev(f"{s['name']} immediate parent", rel, 1, f"Immediate parent: {a['immediate_parent']}"), BODY),
              Paragraph(ev(f"{s['name']} ultimate parent", rel, 1, f"Ultimate parent: {a['ultimate_parent']}"), BODY),
              Spacer(1, 6), Paragraph("Coverage", H2),
              Paragraph(ev(f"{s['name']} coverage period", rel, 1,
                           f"Coverage period: materials delivered to Volterra from {a['coverage'][0]} through {a['coverage'][1]}"), BODY)]
        reps = [
            "The materials listed above were manufactured at the manufacturing site(s) listed above.",
            "The ownership information stated above is complete and accurate as of the date of signature.",
            "Supplier is not a specified foreign entity or a foreign-influenced entity as those terms are defined in IRC section 7701(a)(51).",
            "Supplier will notify Volterra in writing within 30 days of any change to the information above.",
        ]
        if a["ultimate_parent"].startswith("Not disclosed"):
            reps[1] = "Ownership information above the immediate parent is confidential and is not provided in this declaration."
        p2 = [Paragraph("Representations", H2)] + [Paragraph(f"{i + 1}. {r}", BODY) for i, r in enumerate(reps)] + [
            Spacer(1, 12), Paragraph("Signature", H2),
            Paragraph(ev(f"{s['name']} penalty of perjury", rel, 2,
                         "Signed under penalty of perjury by an authorized officer of the supplier."), BODY),
            kv("Name and title", a["signatory"]), kv("Date signed", a["signed"])]
        build_pdf(OUT / rel, s["name"], a["file"].replace(".pdf", ""), [p1, p2])

    for x in EXTRA_SUPPLIER_DOCS:
        s = SUP[x["supplier"]]
        rel = f"SUPPLIER_EVIDENCE/{x['file']}"
        m = MAT_BY_ERP[s["materials"][0]]
        p1 = [Paragraph("Certificate of Origin", H1), kv("Exporter", f"{s['name']}, {s['street']}, {s['city']}, {s['region']} {s['postal']}, Canada"),
              kv("Consignee", "Volterra Battery Systems, Inc., 4100 Lakeshore Industrial Drive, Holland, MI 49423, USA"),
              kv("Goods", f"{m['description']} (Volterra material {m['erp_material']}), HS 3801.10"),
              Paragraph(ev("Kestrel graphite country of origin", rel, 1, "Country of origin of goods: Canada"), BODY),
              kv("Blanket period", "2026-01-01 to 2026-12-31"),
              kv("Certified by", "Marie-Claude Tremblay, VP Finance - 2025-12-18")]
        build_pdf(OUT / rel, s["name"], x["file"].replace(".pdf", ""), [p1])


# ======================================================================================
# 5. DOCUMENTATION
# ======================================================================================

def money(x: float) -> str:
    return f"${x:,.2f}"


def write_docs(orders, lots, consumption, sales, exp):
    internal = OUT / "INTERNAL_TEST_ONLY"
    internal.mkdir(parents=True, exist_ok=True)
    first = orders[0]
    first_sale = next(s for s in sales if s["batch"] == first["fg_lot"])
    lfp = PRODUCTS["VX-LFP50"]
    vx = PRODUCTS["VX-2170"]
    c = exp["credit"]
    prodn = exp["production"]
    sa = exp["sales"]
    mo7 = next(o for o in orders if o["reported_good"] != o["true_good"])
    nordvik_end = ATTESTATIONS["0000100273"]["coverage"][1]
    lbi = exp["lots_by_issue"]
    bg = exp["excluded_unmatched_lines"][0]

    issues_detail = [
        {"id": "ISS-01", "title": ISSUES["ISS-01"], "category": "Identity reconciliation",
         "planted": "Each product has a different identifier in each system: " + "; ".join(
             f"{pid}: PLM {pid} / ERP {p['erp_material']} / MES {p['mes_item']}" for pid, p in PRODUCTS.items())
         + f". The plant is ERP plant {FACILITY['erp_plant_code']} and MES site {FACILITY['mes_site_code']}.",
         "why_realistic": "PLM, ERP and MES are administered by different teams and assign their own keys.",
         "truth": "The identifiers refer to the same two products and the same plant.",
         "customer_evidence": "PLM/product_master.xlsx ERP_MATERIAL_NO; MES/item_master.csv plm_part_ref; ERP descriptions contain the PLM "
                              "number; COMPANY/company_and_facility.xlsx Facilities sheet maps ERP_PLANT to MES_SITE.",
         "expected_behavior": "Reconcile each product and the plant to a single identity using the cross-references in the customer files, "
                              "and show the mapping with its source fields.",
         "expected_status": "RESOLVED_BY_CROSS_REFERENCE", "human_review": False, "affected": list(PRODUCTS)},
        {"id": "ISS-02", "title": ISSUES["ISS-02"], "category": "Supplier ownership - missing evidence",
         "planted": "SUPPLIER_EVIDENCE/EastbaySeparator_Supplier_Declaration_2026.pdf page 1 states 'Ultimate parent: Not disclosed - "
                    "confidential'; ERP supplier master notes only the Singapore immediate parent.",
         "why_realistic": "Foreign suppliers commonly disclose the immediate holding company but decline to identify beneficial owners.",
         "truth": f"Ultimate parent is {SUP['0000100268']['true_ultimate_parent']} ({SUP['0000100268']['true_ultimate_parent_country']}), "
                  f"{SUP['0000100268']['true_ownership_pct']}% indirect ownership. This appears ONLY in ground_truth.json.",
         "customer_evidence": "Immediate parent Eastbay Holdings Pte. Ltd. (Singapore). Ultimate parent not disclosed.",
         "expected_behavior": "Mark the ultimate-parent fact for Eastbay Separator as MISSING EVIDENCE and flag the lots that consumed "
                              "MAT-30000415 for review. Do not state, guess or infer the ultimate parent.",
         "expected_status": "MISSING_EVIDENCE_REQUIRES_REVIEW", "human_review": True, "affected": lbi["ISS-02"]},
        {"id": "ISS-03", "title": ISSUES["ISS-03"], "category": "Supplier evidence - expired",
         "planted": f"SUPPLIER_EVIDENCE/NordvikElectrolyte_Supplier_Attestation_2025-2026.pdf covers materials delivered through {nordvik_end}. "
                    "No later Nordvik attestation exists in the customer files.",
         "why_realistic": "Annual supplier attestations lapse when procurement does not chase renewals.",
         "truth": ATTESTATIONS["0000100273"]["true_note"],
         "customer_evidence": f"Electrolyte lots in MES/material_consumption.csv with component_lot_received after {nordvik_end} are not "
                              "covered by any attestation on file.",
         "expected_behavior": f"Mark electrolyte supplier evidence as EXPIRED for finished lots that consumed MAT-30000520 lots received after "
                              f"{nordvik_end}, and flag those lots for review. Lots using electrolyte received on or before {nordvik_end} are supported. "
                              "Do not assume a renewal exists.",
         "expected_status": "EXPIRED_EVIDENCE_REQUIRES_REVIEW", "human_review": True, "affected": lbi["ISS-03"]},
        {"id": "ISS-04", "title": ISSUES["ISS-04"], "category": "Effective-dated BOM",
         "planted": "PLM/bom.xlsx and PLM/product_master.xlsx contain VX-2170 revision B (2026-01-15 to 2026-06-30) and revision C "
                    "(from 2026-07-01, ECN-2026-031). Revision C replaces RM-SEP-016 (Tamarack) with RM-SEP-021 (Eastbay).",
         "why_realistic": "Engineering change notices switch suppliers mid-year for cost reasons.",
         "truth": "Lots started on or after 2026-07-01 were built to revision C; MES consumption confirms MAT-30000415 was used.",
         "customer_evidence": "BOM effectivity dates; MES consumption of MAT-30000415 only on lots started on or after 2026-07-01.",
         "expected_behavior": "Assign each VX-2170 lot to the BOM revision in effect on its production start date and evaluate supplier "
                              "evidence against that revision's materials.",
         "expected_status": "RESOLVED_BY_EFFECTIVITY", "human_review": False,
         "affected": [o["fg_lot"] for o in orders if o["product"] == "VX-2170"]},
        {"id": "ISS-05", "title": ISSUES["ISS-05"], "category": "Product evidence - missing",
         "planted": "QUALITY/ contains VX2170_RevB_Qualification_Report.pdf only. Page 1 states the report applies only to revision B.",
         "why_realistic": "Re-qualification after a material change is often scheduled but not completed before production continues.",
         "truth": "Revision C cells have the same rated capacity (5.0 Ah); they were simply never re-tested and documented.",
         "customer_evidence": "No test report for VX-2170 revision C. PLM still lists 5.0 Ah for revision C.",
         "expected_behavior": "Mark capacity and specification evidence for VX-2170 revision C lots as MISSING and flag for review. The "
                              "estimate may use the PLM value 5.0 Ah, labelled as not supported by a test report.",
         "expected_status": "MISSING_EVIDENCE_REQUIRES_REVIEW", "human_review": True, "affected": lbi["ISS-05"]},
        {"id": "ISS-06", "title": ISSUES["ISS-06"], "category": "Conflicting evidence",
         "planted": f"PLM/product_master.xlsx RATED_CAPACITY_AH = {lfp['plm_rated_capacity_ah']}; "
                    f"QUALITY/VXLFP50_RevA_Qualification_Report_DRAFT.pdf page 3 'Declared rated capacity: {lfp['true_rated_capacity_ah']} Ah'. "
                    "The report is a draft with no reviewer signature.",
         "why_realistic": "PLM attributes often hold the design target and are not updated after test results.",
         "truth": f"True rated capacity is {lfp['true_rated_capacity_ah']} Ah; the PLM value is a stale design target.",
         "customer_evidence": "Two sources disagree; the only test evidence is an unapproved draft.",
         "expected_behavior": f"Mark VX-LFP50 rated capacity as CONFLICTING, show the credit under both values "
                              f"({money(c['VX-LFP50']['low'])} to {money(c['VX-LFP50']['high'])}), mark neither as supported, and flag for review.",
         "expected_status": "CONFLICTING_EVIDENCE_REQUIRES_REVIEW", "human_review": True, "affected": lbi["ISS-06"]},
        {"id": "ISS-07", "title": ISSUES["ISS-07"], "category": "Quantity discrepancy",
         "planted": f"MES/production_orders.csv {mo7['mes_order']} qty_good = {mo7['reported_good']}; MES/material_consumption.csv "
                    f"fg_lot_qty for lot {mo7['fg_lot']} = {mo7['true_good']}, and material consumed corresponds to {mo7['true_good']} good + "
                    f"{mo7['scrap']} scrap cells.",
         "why_realistic": "Order confirmations are sometimes double-posted after rework while the lot record is correct.",
         "truth": f"{mo7['true_good']} good cells; 62 reworked cells were confirmed twice.",
         "customer_evidence": "Two MES records disagree by 62 cells; material consumption supports the lower figure.",
         "expected_behavior": f"Flag the {mo7['reported_good'] - mo7['true_good']}-cell discrepancy on {mo7['mes_order']} for review and use the "
                              f"finished-lot quantity ({mo7['true_good']}) for available-to-sell calculations, showing both values.",
         "expected_status": "DISCREPANCY_REQUIRES_REVIEW", "human_review": True, "affected": [mo7["fg_lot"]]},
        {"id": "ISS-08", "title": ISSUES["ISS-08"], "category": "Unmatched / ambiguous record",
         "planted": f"ERP/sales.csv invoice {bg['invoice']}: MATNR {bg['erp_material']} '{BGRADE['description']}', {bg['qty']} EA to "
                    f"{CUSTOMERS[BGRADE['customer']]['name']}, no batch (CHARG blank). {bg['erp_material']} is not in the PLM product master "
                    "or MES item master.",
         "why_realistic": "Off-spec product is sold to recyclers under separate material numbers outside normal lot tracking.",
         "truth": BGRADE["true_origin"],
         "customer_evidence": "Description resembles VX-LFP50 but no identifier or lot links it to an eligible product.",
         "expected_behavior": "Exclude the line from the credit estimate and flag it for review as an unmatched sales record. Do not map "
                              "it to VX-LFP50 based on the description.",
         "expected_status": "EXCLUDED_UNMATCHED_REQUIRES_REVIEW", "human_review": True, "affected": [bg["invoice"]]},
    ]

    # ---- ground_truth.json
    gt = {
        "meta": {"description": "PRIVATE test oracle for the Volterra synthetic dataset. Not a customer-provided file.",
                 "tax_year": TAX_YEAR, "generated_by": "generate_dataset.py", "all_entities_fictional": True},
        "claimant": COMPANY, "facility": FACILITY,
        "products": {pid: {**{k: v for k, v in p.items()}, "volume_l": round(cell_volume_l(p), 5)} for pid, p in PRODUCTS.items()},
        "identifier_mappings": {pid: {"plm": pid, "erp_material": p["erp_material"], "mes_item": p["mes_item"]} for pid, p in PRODUCTS.items()}
                               | {"_facility": {"erp_plant": FACILITY["erp_plant_code"], "mes_site": FACILITY["mes_site_code"],
                                                "facility_id": FACILITY["facility_id"]}},
        "materials": MATERIALS, "subassemblies": SUBASSEMBLIES,
        "boms": {f"{pid} rev {r['rev']}": bom_lines(pid, r["rev"]) for pid, p in PRODUCTS.items() for r in p["revisions"]},
        "suppliers": SUPPLIERS, "supplier_documents": ATTESTATIONS, "customers": CUSTOMERS,
        "production_orders": [{k: (v.isoformat() if isinstance(v, date) else v) for k, v in o.items()} for o in orders],
        "material_lots": [{"material": l["material"], "lot": l["lot"], "received": l["received"].isoformat(), "supplier": l["supplier"],
                           "size": base_to_erp(l["size_base"], MAT_BY_PLM[l["material"]]["base"])} for ls in lots.values() for l in ls],
        "sales": [{**{k: v for k, v in s.items() if k != "date"}, "date": s["date"].isoformat()} for s in sales],
        "bgrade_truth": BGRADE,
        "evidence_index": EVIDENCE_INDEX,
        "intentional_issues": issues_detail,
        "tax_rules_applied": {k: v["citation"] for k, v in tr.RULES.items()},
        "assumptions": tr.ASSUMPTIONS, "unresolved_rules": tr.UNRESOLVED_RULES,
        "expected": exp,
    }
    (internal / "ground_truth.json").write_text(json.dumps(gt, indent=2, default=str) + "\n", encoding="utf-8")

    # ---- intentional_issues.md
    md = ["# Intentional issues (INTERNAL - TEST ONLY)", "",
          "Every deliberately introduced discrepancy or gap in the customer-facing dataset. There are exactly "
          f"{len(issues_detail)}; `validate_dataset.py` checks that the files contain these and no others.", "",
          "Each issue has exactly one expected Marigold behavior. Marigold must never use information that exists only in "
          "`ground_truth.json`; where evidence is missing or conflicting, the expected behavior is to flag for review.", ""]
    for i in issues_detail:
        md += [f"## {i['id']} - {i['title']}", "",
               f"- **Category:** {i['category']}",
               f"- **What was planted:** {i['planted']}",
               f"- **Why it is realistic:** {i['why_realistic']}",
               f"- **Private truth (ground_truth.json only):** {i['truth']}",
               f"- **What the customer-facing evidence shows:** {i['customer_evidence']}",
               f"- **Expected Marigold behavior (the one correct behavior):** {i['expected_behavior']}",
               f"- **Expected status:** `{i['expected_status']}`",
               f"- **Human review expected:** {'Yes' if i['human_review'] else 'No'}",
               f"- **Affected records:** {', '.join(i['affected'])}", ""]
    (internal / "intentional_issues.md").write_text("\n".join(md), encoding="utf-8")

    # ---- expected_results.md
    pf = exp["product_facts"]
    vx_fact = pf["VX-2170"]["by_capacity_value"][str(vx["true_rated_capacity_ah"])]
    lf_hi = pf["VX-LFP50"]["by_capacity_value"][str(lfp["plm_rated_capacity_ah"])]
    lf_lo = pf["VX-LFP50"]["by_capacity_value"][str(lfp["true_rated_capacity_ah"])]
    vxc = c["VX-2170"]["by_capacity_basis"]["plm_value"]
    lfc = c["VX-LFP50"]["by_capacity_basis"]
    order_rows = "\n".join(
        f"| {o['mes_order']} | {o['fg_lot']} | {o['product']} | {o['bom_rev']} | {o['start']} | {o['end']} | {o['planned']:,} | "
        f"{o['reported_good']:,} | {o['true_good']:,} | {o['scrap']:,} | {', '.join(exp['lot_open_issues'][o['fg_lot']]) or 'none'} |"
        for o in orders)
    md = f"""# Expected results (INTERNAL - TEST ONLY)

What a correctly functioning Marigold prototype should conclude from the customer-facing files for tax year {TAX_YEAR}.
All numbers below were computed by `generate_dataset.py`; `validate_dataset.py` recomputes them from the files.

Three layers are kept separate throughout:

1. **True** - what `ground_truth.json` says actually happened.
2. **Established** - what can be shown from the customer-facing files.
3. **Unresolved** - what the customer files cannot establish. Marigold must flag these, not resolve them from the ground truth.

Status vocabulary: `SUPPORTED`, `RESOLVED_BY_CROSS_REFERENCE`, `RESOLVED_BY_EFFECTIVITY`, `MISSING_EVIDENCE_REQUIRES_REVIEW`,
`EXPIRED_EVIDENCE_REQUIRES_REVIEW`, `CONFLICTING_EVIDENCE_REQUIRES_REVIEW`, `DISCREPANCY_REQUIRES_REVIEW`,
`EXCLUDED_UNMATCHED_REQUIRES_REVIEW`, `NOT_COMPUTED_REQUIRES_REVIEW`.

## 1. Summary

| Conclusion | Expected value | Status |
|---|---|---|
| Claimant | {COMPANY['legal_name']} | SUPPORTED |
| Facility | {FACILITY['name']}, {FACILITY['city']}, {FACILITY['state']} (United States) | SUPPORTED |
| Claimed component type | Battery cells only (VX-2170, VX-LFP50) | SUPPORTED |
| VX-2170 cells sold in {TAX_YEAR} | {sa['VX-2170']['qty_2026']:,} | SUPPORTED |
| VX-LFP50 cells sold in {TAX_YEAR} | {sa['VX-LFP50']['qty_2026']:,} | SUPPORTED |
| VX-2170 estimated credit | {money(c['VX-2170']['high'])} | Partly SUPPORTED, partly requires review |
| VX-LFP50 estimated credit | {money(c['VX-LFP50']['low'])} to {money(c['VX-LFP50']['high'])} | CONFLICTING_EVIDENCE_REQUIRES_REVIEW |
| **Total estimated credit** | **{money(exp['credit_total']['low'])} to {money(exp['credit_total']['high'])}** | Requires review |
| Credit on lots with no open issues | {money(exp['credit_total']['no_open_issues'])} | SUPPORTED |
| Material-assistance (PFE) test | Inputs provided, no pass/fail | NOT_COMPUTED_REQUIRES_REVIEW |

## 2. Claimant and facility

- **Conclusion:** {COMPANY['legal_name']} is the claimant; it owns and operates the only facility, {FACILITY['name']}, {FACILITY['street']}, {FACILITY['city']}, {FACILITY['state']} {FACILITY['postal_code']}, United States.
- **Status:** SUPPORTED. **Human review:** No.
- **Evidence:** `COMPANY/company_and_facility.xlsx` (LegalEntity and Facilities sheets); MES `site` = {FACILITY['mes_site_code']} and ERP `WERKS` = {FACILITY['erp_plant_code']} map to {FACILITY['facility_id']} (ISS-01).
- **Related parties:** the LegalEntity sheet lists no parent, subsidiaries or affiliates, and none of the four customers is a Volterra entity, so all sales are to unrelated persons (SUPPORTED).

## 3. Product identity (ISS-01)

| PLM | ERP material | MES item | Reconciled via | Status |
|---|---|---|---|---|
""" + "\n".join(f"| {pid} | {p['erp_material']} | {p['mes_item']} | product_master ERP_MATERIAL_NO; item_master plm_part_ref | RESOLVED_BY_CROSS_REFERENCE |"
                for pid, p in PRODUCTS.items()) + f"""

`{BGRADE['erp_material']}` ({BGRADE['description']}) appears only in `ERP/sales.csv` and does not reconcile to any product (ISS-08).

## 4. Preliminary eligibility by product

Battery-cell definition tests (Form 7207 instructions, line 5a; see `tax_rules.py`): at least {tr.RULES['BATTERY_CELL_MIN_ENERGY_WH']['value']:.0f} Wh, at least {tr.RULES['BATTERY_CELL_MIN_ENERGY_DENSITY_WH_PER_L']['value']:.0f} Wh/L, capacity-to-power ratio at most {tr.RULES['BATTERY_CELL_MAX_CAPACITY_TO_POWER_RATIO']['value']:.0f}:1.

### VX-2170 (cylindrical NMC)

| Fact | Value | Evidence | Status |
|---|---|---|---|
| Nominal voltage | {vx['nominal_voltage_v']} V | QUALITY/VX2170_RevB_Qualification_Report.pdf p.2 | SUPPORTED (rev B) |
| Rated capacity | {vx['true_rated_capacity_ah']} Ah | same report p.3; PLM RATED_CAPACITY_AH | SUPPORTED for rev B; MISSING for rev C (ISS-05) |
| Energy per cell | {vx_fact['energy_wh']} Wh (>= 12 Wh) | derived | SUPPORTED (rev B) |
| Volumetric energy density | {vx_fact['energy_density_wh_per_l']} Wh/L (>= 100) | report p.3 (external volume p.2) | SUPPORTED (rev B) |
| Capacity-to-power ratio | {vx_fact['capacity_to_power_ratio']} : 1 (<= 100) | report p.4 | SUPPORTED (rev B) |
| Produced in the U.S. | Holland, MI | MES site HOLLAND-01 -> company file | SUPPORTED |

- **Conclusion:** meets the battery-cell definition on revision B evidence. Revision C lots ({', '.join(lbi['ISS-05'])}) lack a qualification report.
- **Expected status:** revision B lots SUPPORTED; revision C lots MISSING_EVIDENCE_REQUIRES_REVIEW (ISS-05), plus supplier issues ISS-02 and ISS-03.

### VX-LFP50 (prismatic LFP)

| Fact | PLM value | Test report value (draft) | Status |
|---|---|---|---|
| Rated capacity | {lfp['plm_rated_capacity_ah']} Ah | {lfp['true_rated_capacity_ah']} Ah (p.3) | CONFLICTING_EVIDENCE_REQUIRES_REVIEW (ISS-06) |
| Energy per cell | {lf_hi['energy_wh']} Wh | {lf_lo['energy_wh']} Wh | both >= 12 Wh |
| Volumetric energy density | {lf_hi['energy_density_wh_per_l']} Wh/L | {lf_lo['energy_density_wh_per_l']} Wh/L | both >= 100 Wh/L |
| Capacity-to-power ratio | {lf_hi['capacity_to_power_ratio']} : 1 | {lf_lo['capacity_to_power_ratio']} : 1 | both <= 100:1 |

- **Conclusion:** meets the battery-cell definition under either capacity value, but the credit amount depends on which value is correct, and the only test evidence is an unapproved draft.
- **Expected status:** CONFLICTING_EVIDENCE_REQUIRES_REVIEW for all VX-LFP50 lots. **Human review:** Yes.

## 5. Production ({TAX_YEAR})

| MES order | Lot | Product | BOM rev | Start | End | Planned | MES qty_good | Finished-lot qty | Scrap | Open issues on lot |
|---|---|---|---|---|---|---|---|---|---|---|
{order_rows}

| Product | Orders | Planned | MES reported good | Finished-lot good (use this) | Scrap |
|---|---|---|---|---|---|
""" + "\n".join(f"| {pid} | {v['orders']} | {v['planned']:,} | {v['mes_reported_good']:,} | {v['finished_lot_good']:,} | {v['scrap']:,} |"
                for pid, v in prodn.items()) + f"""

- **Known discrepancy:** {mo7['mes_order']} reports {mo7['reported_good']:,} good cells, the finished lot shows {mo7['true_good']:,} (ISS-07).

## 6. Sales ({TAX_YEAR})

| Product | 2026 invoices | 2026 lines | Cells sold in 2026 | Revenue 2026 | Ending inventory 12/31/2026 | Jan 2027 cells (excluded) |
|---|---|---|---|---|---|---|
""" + "\n".join(f"| {pid} | {v['invoices_2026']} | {v['lines_2026']} | {v['qty_2026']:,} | {money(v['revenue_2026'])} | {v['ending_inventory_2026_12_31']:,} | {v['qty_excluded_2027']:,} (invoices {', '.join(v['invoices_excluded_2027'])}) |"
                for pid, v in sa.items()) + f"""

- `ERP/sales.csv` contains {exp['sales_counts']['invoices_total']} invoices ({exp['sales_counts']['lines_total']} lines), of which {exp['sales_counts']['invoices_2026']} invoices ({exp['sales_counts']['lines_2026']} lines) are dated {TAX_YEAR}.
- Invoices dated January 2027 are outside the tax year and must be excluded (scope rule, not an intentional issue).
- Invoice {bg['invoice']} ({bg['qty']} x {bg['erp_material']}, {money(bg['amount'])}) is excluded and flagged (ISS-08).
- Cumulative cells sold never exceed cumulative finished-lot output available before each invoice date (checked by the validator).

## 7. Credit calculation

Rule inputs (`tax_rules.py`): ${tr.RULES['BATTERY_CELL_RATE_PER_KWH']['value']:.2f}/kWh for battery cells, phase-out factor {tr.RULES['PHASE_OUT_FACTOR']['value'][TAX_YEAR]:.0%} for {TAX_YEAR} sales, kWh per cell = nominal V x rated Ah / 1000 (assumption A1).

| Product | Capacity basis | Rated Ah | kWh per cell | Cells sold 2026 | Total kWh | Estimated credit | On lots with no open issues | On lots with open issues |
|---|---|---|---|---|---|---|---|---|
| VX-2170 | PLM = test report | {vxc['rated_capacity_ah']} | {vxc['kwh_per_cell']} | {c['VX-2170']['units_sold_2026']:,} | {vxc['kwh_total']:,} | {money(vxc['credit'])} | {money(vxc['credit_no_open_issues'])} | {money(vxc['credit_with_open_issues'])} |
| VX-LFP50 | PLM value | {lfc['plm_value']['rated_capacity_ah']} | {lfc['plm_value']['kwh_per_cell']} | {c['VX-LFP50']['units_sold_2026']:,} | {lfc['plm_value']['kwh_total']:,} | {money(lfc['plm_value']['credit'])} | {money(lfc['plm_value']['credit_no_open_issues'])} | {money(lfc['plm_value']['credit_with_open_issues'])} |
| VX-LFP50 | Draft test report | {lfc['test_report_value']['rated_capacity_ah']} | {lfc['test_report_value']['kwh_per_cell']} | {c['VX-LFP50']['units_sold_2026']:,} | {lfc['test_report_value']['kwh_total']:,} | {money(lfc['test_report_value']['credit'])} | {money(lfc['test_report_value']['credit_no_open_issues'])} | {money(lfc['test_report_value']['credit_with_open_issues'])} |

- **Total estimated credit:** {money(exp['credit_total']['low'])} (test-report basis) to {money(exp['credit_total']['high'])} (PLM basis). Status: requires review.
- **Credit with no open issues:** {money(exp['credit_total']['no_open_issues'])} - VX-2170 sales allocated to lots with no open issues ({', '.join(o['fg_lot'] for o in orders if not exp['lot_open_issues'][o['fg_lot']])}).
- **Credit touched by each open issue (PLM basis, overlapping):** """ + "; ".join(f"{k} {money(v)}" for k, v in sorted(exp["credit_touched_by_issue_plm_basis"].items())) + f"""

## 8. Evidence and substantiation status by lot

| Lot | Product | BOM rev | Open issues | Lot status |
|---|---|---|---|---|
""" + "\n".join(f"| {o['fg_lot']} | {o['product']} | {o['bom_rev']} | {', '.join(exp['lot_open_issues'][o['fg_lot']]) or 'none'} | "
                f"{'SUPPORTED' if not exp['lot_open_issues'][o['fg_lot']] else 'REQUIRES REVIEW'} |" for o in orders) + f"""

## 9. Supplier issues

| Supplier | Material | Evidence on file | Expected status |
|---|---|---|---|
""" + "\n".join(
        f"| {s['name']} | {', '.join(s['materials'])} | {ATTESTATIONS[s['erp_id']]['file']} (coverage {ATTESTATIONS[s['erp_id']]['coverage'][0]} to {ATTESTATIONS[s['erp_id']]['coverage'][1]}) | "
        + ("MISSING_EVIDENCE_REQUIRES_REVIEW - ultimate parent not disclosed (ISS-02)" if s["erp_id"] == "0000100268"
           else f"EXPIRED_EVIDENCE_REQUIRES_REVIEW for lots received after {nordvik_end} (ISS-03)" if s["erp_id"] == "0000100273"
           else "SUPPORTED") + " |" for s in SUPPLIERS) + f"""

Electrode active materials (NMC811 and LFP cathode material, synthetic graphite) are purchased inputs. Volterra does not claim a credit for them.

## 10. Material-assistance inputs (not computed)

Status: NOT_COMPUTED_REQUIRES_REVIEW (`tax_rules.UNRESOLVED_RULES` U1). Inputs available from the customer files:

| BOM | Direct material cost per cell (purchased materials) |
|---|---|
""" + "\n".join(f"| {k} | {money(v['direct_material_cost_per_cell'])} |" for k, v in exp["material_assistance_inputs"]["cost_per_cell_by_bom_revision"].items()) + """

| Supplier | Cost of material consumed in 2026 |
|---|---|
""" + "\n".join(f"| {k} | {money(v)} |" for k, v in exp["material_assistance_inputs"]["consumed_material_cost_2026_by_supplier"].items()) + """

## 11. Unresolved issues (must remain unresolved)

""" + "\n".join(f"- **{i['id']}** {i['title']} - expected status `{i['expected_status']}`. {i['expected_behavior']}"
                for i in issues_detail if i["human_review"]) + "\n"
    (internal / "expected_results.md").write_text(md, encoding="utf-8")

    write_data_dictionary(internal)
    write_readme(first, first_sale, consumption)


DICT = [
    ("COMPANY/company_and_facility.xlsx", "Corporate tax/legal records (maintained by Tax)",
     "Legal entity, synthetic EIN, tax year, facility address and identifiers, production lines.",
     "LegalEntity: FIELD. Facilities: FACILITY_ID. ProductionLines: (FACILITY_ID, ERP_WORK_CENTER).",
     "Facilities.ERP_PLANT -> ERP/sales.csv WERKS; Facilities.MES_SITE -> MES/production_orders.csv site; ProductionLines.MES_LINE -> production_orders.line.",
     [("LegalEntity.FIELD / VALUE", "text", "", "Key-value facts about the legal entity, including related entities."),
      ("Facilities.FACILITY_ID", "text", "", "Corporate facility identifier."),
      ("Facilities.ERP_PLANT", "text", "", "Plant code used by ERP (SAP-style WERKS)."),
      ("Facilities.MES_SITE", "text", "", "Site code used by MES."),
      ("Facilities.STREET/CITY/STATE/ZIP/COUNTRY", "text", "", "Physical address."),
      ("Facilities.OWNER / OPERATOR", "text", "", "Entity owning / operating the plant."),
      ("Facilities.OPERATIONS_START", "date", "", "First day of operations."),
      ("ProductionLines.ERP_WORK_CENTER / MES_LINE / DESCRIPTION", "text", "", "Line identifiers in each system.")],
     "Claimant identity; facility is in the United States; plant identifier reconciliation; related-party status."),
    ("PLM/product_master.xlsx", "PLM (product lifecycle management) export",
     "One row per product revision with released attributes.", "(PART_NUMBER, REVISION)", "ERP_MATERIAL_NO -> ERP/sales.csv MATNR; PART_NUMBER -> MES/item_master.csv plm_part_ref; PART_NUMBER -> bom.xlsx TOP_ASSEMBLY.",
     [("PART_NUMBER", "text", "", "PLM part number of the finished cell."), ("REVISION", "text", "", "Engineering revision."),
      ("PART_NAME / DESCRIPTION", "text", "", "Names."), ("LIFECYCLE_STATE", "text", "", "Released or Superseded."),
      ("EFFECTIVE_FROM / EFFECTIVE_TO", "date", "", "Revision effectivity (blank TO = open-ended)."),
      ("CHANGE_NOTICE / CHANGE_DESCRIPTION", "text", "", "Engineering change notice that released the revision."),
      ("ERP_MATERIAL_NO", "text", "", "Cross-reference to the ERP material number."), ("ITEM_CLASS", "text", "", "PLM item class."),
      ("FORM_FACTOR / CHEMISTRY", "text", "", "Cell construction."), ("NOMINAL_VOLTAGE_V", "number", "V", "Nominal voltage."),
      ("RATED_CAPACITY_AH", "number", "Ah", "Rated capacity attribute held in PLM."), ("DIMENSIONS_MM", "text", "mm", "Nominal external dimensions."),
      ("MASS_G", "number", "g", "Nominal mass."), ("MAX_CONT_DISCHARGE_A", "number", "A", "Maximum continuous discharge current."),
      ("CHARGE_VOLTAGE_V / DISCHARGE_CUTOFF_V", "number", "V", "Voltage limits."), ("SPEC_DOC", "text", "", "Engineering specification reference.")],
     "Product classification inputs (voltage, capacity, dimensions); revision effectivity; identifier mapping."),
    ("PLM/bom.xlsx", "PLM bill of materials export",
     "Two-level product structure per revision: cell -> electrode subassemblies (made in-house) -> purchased materials.",
     "(TOP_ASSEMBLY, TOP_REV, FIND_NO)", "COMPONENT_ITEM (purchased) -> ERP_MATERIAL_NO -> ERP/supplier_master MATERIALS_SUPPLIED and material_purchase_prices MATNR; PARENT_ITEM -> COMPONENT_ITEM of the level above.",
     [("TOP_ASSEMBLY / TOP_REV", "text", "", "Finished cell and revision the line belongs to."), ("BOM_LEVEL", "integer", "", "1 = direct child of the cell, 2 = child of a subassembly."),
      ("FIND_NO", "text", "", "Position number."), ("PARENT_ITEM", "text", "", "Immediate parent of the component."),
      ("COMPONENT_ITEM / COMPONENT_DESCRIPTION", "text", "", "PLM component part number and description."),
      ("QTY_PER", "number", "per UOM", "Quantity per one parent item."), ("UOM", "text", "", "G (grams), M2 (square meters) or EA (each)."),
      ("MAKE_BUY", "text", "", "Make = produced in-house; Purchased = bought from a supplier."), ("ERP_MATERIAL_NO", "text", "", "ERP material for purchased components."),
      ("EFFECTIVE_FROM / EFFECTIVE_TO / CHANGE_NOTICE", "date/text", "", "Effectivity of the revision.")],
     "What goes into each cell, by revision; which purchased materials (and therefore suppliers) affect which lots."),
    ("MES/item_master.csv", "MES (manufacturing execution system) item master",
     "MES item codes for finished cells.", "item_code", "plm_part_ref -> PLM/product_master.xlsx PART_NUMBER; item_code -> production_orders.csv item_code.",
     [("item_code", "text", "", "MES item code."), ("item_description", "text", "", "MES description."), ("plm_part_ref", "text", "", "PLM part number reference."),
      ("routing_version", "text", "", "MES routing version (not the PLM revision)."), ("default_line", "text", "", "MES line code."), ("uom", "text", "", "Unit of measure."), ("status", "text", "", "Item status.")],
     "Identifier mapping between MES and PLM."),
    ("MES/production_orders.csv", "MES production order history",
     "One row per production order; each order produces one finished lot.", "order_id", "item_code -> item_master; site -> company Facilities.MES_SITE; lot_id -> material_consumption.fg_lot_id and ERP/sales.csv CHARG; erp_order_ref = ERP process order.",
     [("order_id", "text", "", "MES order number."), ("erp_order_ref", "text", "", "ERP process order number."), ("item_code", "text", "", "MES item code."),
      ("site / line", "text", "", "MES site and line codes."), ("lot_id", "text", "", "Finished-goods lot number."), ("start_ts / end_ts", "ISO timestamp", "local time", "Order start/end."),
      ("qty_planned", "integer", "EA", "Planned quantity."), ("qty_good", "integer", "EA", "Good quantity confirmed on the order."), ("qty_scrap", "integer", "EA", "Scrapped cells."),
      ("uom", "text", "", "EA."), ("order_status", "text", "", "Order status.")],
     "Production quantities and dates; U.S. production (via site); lots available to sell."),
    ("MES/material_consumption.csv", "MES lot genealogy / material consumption",
     "What purchased material lots were actually consumed by each order (backflushed through the electrode subassemblies).",
     "(order_id, component_material, component_lot)", "order_id -> production_orders; fg_lot_id -> production_orders.lot_id; component_material -> ERP material / PLM BOM ERP_MATERIAL_NO.",
     [("order_id", "text", "", "MES order."), ("fg_lot_id", "text", "", "Finished lot."), ("fg_lot_qty", "integer", "EA", "Good quantity recorded on the finished lot at lot close."),
      ("component_material", "text", "", "ERP material number consumed."), ("component_lot", "text", "", "Supplier lot consumed."),
      ("component_lot_received", "date", "", "Date the supplier lot was received at the plant."), ("qty_consumed", "number", "KG / M2 / EA", "Quantity consumed (good + scrapped cells)."),
      ("uom", "text", "", "KG, M2 or EA (note: BOM uses G and M2)."), ("posted_ts", "ISO timestamp", "", "Posting time.")],
     "Actual material usage by lot; which suppliers' material went into which lots; delivery dates for evidence coverage; finished-lot quantities."),
    ("ERP/supplier_master.xlsx", "ERP vendor master (SAP-style field names)",
     "One row per supplier.", "LIFNR", "MATERIALS_SUPPLIED -> ERP material numbers; LIFNR -> material_purchase_prices.LIFNR.",
     [("LIFNR", "text", "", "Vendor number."), ("NAME1", "text", "", "Supplier legal name."), ("LAND1", "text", "", "Country (ISO 3166 alpha-2)."),
      ("STRAS / ORT01 / REGIO / PSTLZ", "text", "", "Street, city, region, postal code."), ("MATERIALS_SUPPLIED", "text", "", "Materials bought from this vendor."),
      ("PARENT_COMPANY_NOTE", "text", "", "Free-text parent information recorded by procurement at onboarding (not verified)."),
      ("ZTERM", "text", "", "Payment terms."), ("STATUS", "text", "", "Vendor status.")],
     "Supplier identity, country and materials supplied; procurement's (unverified) parent note."),
    ("ERP/material_purchase_prices.xlsx", "ERP purchasing info records / price list",
     "2026 purchase price per material.", "MATNR", "LIFNR -> supplier_master.",
     [("MATNR / MAKTX", "text", "", "Material number and description."), ("MEINS", "text", "", "Purchasing unit (KG, M2, EA)."), ("NETPR", "number", "USD per MEINS", "Net price."),
      ("WAERS", "text", "", "Currency."), ("LIFNR / VENDOR_NAME", "text", "", "Vendor."), ("VALID_FROM / VALID_TO", "date", "", "Price validity.")],
     "Direct material costs (inputs to a future material-assistance calculation)."),
    ("ERP/sales.csv", "ERP billing document export (SAP-style field names)",
     "One row per invoice line, January 2026 to mid-January 2027.", "(VBELN, POSNR)", "MATNR -> PLM product_master ERP_MATERIAL_NO; CHARG -> MES production_orders.lot_id; WERKS -> company Facilities.ERP_PLANT.",
     [("VBELN", "text", "", "Invoice number."), ("POSNR", "text", "", "Invoice line."), ("FKDAT", "date MM/DD/YYYY", "", "Invoice date."),
      ("KUNNR / NAME1", "text", "", "Customer number and name."), ("MATNR / ARKTX", "text", "", "Material sold and line text."),
      ("CHARG", "text", "", "Batch (finished lot) shipped; blank if not batch-managed."), ("FKIMG", "integer", "EA", "Billed quantity."),
      ("VRKME", "text", "", "Sales unit."), ("NETPR", "number", "USD/EA", "Unit price."), ("NETWR", "number", "USD", "Net line value."),
      ("WAERK", "text", "", "Currency."), ("WERKS", "text", "", "Shipping plant."), ("LAND1", "text", "", "Ship-to country.")],
     "Units sold, sale dates (tax-year scope), customers, link from sale to production lot."),
    ("QUALITY/*.pdf", "Quality / engineering document repository",
     "Cell qualification test reports (one per tested product revision).", "Document number in header", "Product and revision stated on page 1.",
     [("Page 1", "", "", "Document status, product/revision scope, summary."), ("Page 2", "", "", "Identification, voltage, dimensions, external volume."),
      ("Page 3", "", "", "Capacity samples, declared rated capacity, rated energy, energy density."), ("Page 4", "", "", "Rate capability, capacity-to-power ratio, sign-off.")],
     "Technical specifications for battery-cell classification and kWh capacity; approval status."),
    ("SUPPLIER_EVIDENCE/*.pdf", "Supplier compliance documents (procurement document store)",
     "Supplier attestations, a declaration, a questionnaire and a certificate of origin.", "File name", "Supplier legal name and Volterra vendor number on page 1; materials covered by ERP material number.",
     [("Page 1", "", "", "Supplier identity, materials and manufacturing sites, immediate and ultimate parent, coverage period."),
      ("Page 2", "", "", "Representations and signature (penalty of perjury).")],
     "Supplier ownership, manufacturing origin, representation coverage by delivery date."),
]


def write_data_dictionary(internal: Path) -> None:
    out = ["# Data dictionary (INTERNAL)", "",
           "Every customer-facing file: the fictional source system, what it represents, keys, relationships and the Marigold facts it can support.", ""]
    for path, system, what, pk, fk, cols, facts in DICT:
        out += [f"## `{path}`", "", f"- **Source system:** {system}", f"- **Represents:** {what}", f"- **Primary key:** {pk}",
                f"- **Relationships:** {fk}", f"- **Can support:** {facts}", "", "| Column | Type | Unit | Meaning |", "|---|---|---|---|"]
        out += [f"| {c} | {t} | {u} | {m} |" for c, t, u, m in cols]
        out.append("")
    (internal / "data_dictionary.md").write_text("\n".join(out), encoding="utf-8")


def write_readme(first, first_sale, consumption) -> None:
    cons = [c for c in consumption if c["order"] is first]
    cam = next(c for c in cons if c["material"]["plm_part"] == "RM-CAM-811")
    text = f"""# Volterra Battery Systems, Inc. - data package for 2026

FICTIONAL DATA. Every company, person, address, identifier and transaction in this package is synthetic.

## About Volterra

Volterra Battery Systems, Inc. manufactures lithium-ion battery cells at one plant, the Volterra Holland Cell Plant in
Holland, Michigan. It started production in 2026 and makes two cells:

- **VX-2170** - a 21700 cylindrical NMC cell sold to power-tool and e-bike makers.
- **VX-LFP50** - a 50 Ah-class prismatic LFP cell sold to a stationary-storage integrator.

Volterra buys cathode active material, graphite anode material, separator, electrolyte and cell cans/cases from outside
suppliers, coats its own electrodes and assembles the cells.

## What is in each folder

| Folder | Source system | Contents |
|---|---|---|
| `COMPANY/` | Corporate tax records | Legal entity, tax year, facility address and the plant codes each system uses |
| `PLM/` | Product lifecycle management | Product master (attributes by revision) and bill of materials |
| `MES/` | Manufacturing execution system | Item master, production orders (one finished lot each) and material consumption (lot genealogy) |
| `ERP/` | Enterprise resource planning | Vendor master, purchase price list and invoice lines (Jan 2026 - mid-Jan 2027) |
| `QUALITY/` | Quality document repository | Cell qualification test reports |
| `SUPPLIER_EVIDENCE/` | Procurement document store | Supplier attestations, declarations, questionnaires and certificates |

## Identifiers differ between systems

Each system uses its own keys. They can be reconciled through cross-reference fields:

| Thing | PLM | ERP | MES | Where the link is |
|---|---|---|---|---|
| 21700 cell | VX-2170 | MAT-10004721 | VX2170_R7 | PLM `ERP_MATERIAL_NO`; MES `plm_part_ref` |
| Prismatic cell | VX-LFP50 | MAT-10004755 | LFP50-PR-A | PLM `ERP_MATERIAL_NO`; MES `plm_part_ref` |
| Plant | - | MI01 (`WERKS`) | HOLLAND-01 (`site`) | `COMPANY/company_and_facility.xlsx`, Facilities sheet |
| Purchased materials | RM-... (BOM `COMPONENT_ITEM`) | MAT-3000... | MAT-3000... | BOM `ERP_MATERIAL_NO` |
| Finished lot | - | `CHARG` on invoices | `lot_id` | same value |

Units also differ: the BOM states grams and square meters per cell, while MES records consumption in kilograms and
square meters. MES `routing_version` is not the PLM engineering revision. Dates are ISO in MES, MM/DD/YYYY in ERP
invoices and Excel dates in PLM.

## How the records relate

```
PLM product (revision, effective dates) --BOM--> purchased materials --ERP material--> supplier (vendor master)
      |                                                                                      |
      +--ERP_MATERIAL_NO / plm_part_ref                                    supplier documents (SUPPLIER_EVIDENCE)
      v
MES production order --lot_id--> material consumption (which supplier lots, received when)
      |
      +--lot_id = CHARG--> ERP invoice line (customer, date, quantity)
      |
QUALITY test report (product + revision) --> technical specifications
```

## Worked example: tracing one VX-2170 lot

1. **Product definition** - `PLM/product_master.xlsx`: VX-2170 revision B, effective 2026-01-15, 3.6 V, 5.0 Ah, ERP material MAT-10004721.
2. **Bill of materials** - `PLM/bom.xlsx`: VX-2170 rev B contains RM-CAM-811 (24.5 g NMC811 per cell, ERP MAT-30000102) inside cathode electrode SA-CE-2170.
3. **Supplier** - `ERP/supplier_master.xlsx`: MAT-30000102 is supplied by Great Plains Cathode LLC (vendor 0000100231, Lincoln, NE).
4. **Production** - `MES/production_orders.csv`: order {first['mes_order']} (ERP order {first['erp_order']}) produced lot {first['fg_lot']}, {first['true_good']:,} good cells, {first['start']} to {first['end']} at site HOLLAND-01.
5. **Actual material used** - `MES/material_consumption.csv`: lot {first['fg_lot']} consumed {base_to_erp(cam['qty_base'], 'dg'):.4f} KG of MAT-30000102 from supplier lot {cam['lot']['lot']}, received {cam['lot']['received']}.
6. **Sale** - `ERP/sales.csv`: invoice {first_sale['invoice']} on {first_sale['date']:%m/%d/%Y} shipped {first_sale['qty']:,} cells of MAT-10004721 from batch {first_sale['batch']} to {CUSTOMERS[first_sale['customer']]['name']}.
7. **Supporting evidence** -
   - `QUALITY/VX2170_RevB_Qualification_Report.pdf` (rated capacity page 3, voltage and volume page 2);
   - `SUPPLIER_EVIDENCE/GreatPlainsCathode_Supplier_Attestation_2026.pdf` (manufacturing site, ownership and coverage page 1, signature page 2).
"""
    (OUT / "README.md").write_text(text, encoding="utf-8")


# ======================================================================================

def main() -> None:
    if OUT.exists():
        shutil.rmtree(OUT)
    OUT.mkdir(parents=True)
    orders, lots, consumption, sales, _ = simulate()
    write_company()
    write_plm(orders)
    write_mes(orders, consumption)
    write_erp(sales)
    write_quality()
    write_supplier_docs()
    exp = compute_expected(orders, lots, consumption, sales)
    write_docs(orders, lots, consumption, sales, exp)
    files = sorted(p.relative_to(OUT) for p in OUT.rglob("*") if p.is_file())
    print(f"Generated {len(files)} files in {OUT.relative_to(HERE)}/")
    for f in files:
        print(f"  {f}")
    print(f"Estimated 2026 credit: {money(exp['credit_total']['low'])} - {money(exp['credit_total']['high'])}; "
          f"no open issues: {money(exp['credit_total']['no_open_issues'])}")


if __name__ == "__main__":
    main()
