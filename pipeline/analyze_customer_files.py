#!/usr/bin/env python3
"""Simulated Marigold pipeline: analyze Volterra's customer-facing files and emit the dashboard analysis.

Reads ONLY customer-provided files (never INTERNAL_TEST_ONLY/), applies only the rules in
tax_rules_45x.py (each quoted from docs/compliance/), and writes:

    lib/analysis/dashboard_analysis.json   - structured analysis consumed by the UI
    pipeline/ANALYSIS_SUMMARY.md           - concise internal analysis for reviewers

Run from the repository root:
    data/volterra_synthetic/.venv/bin/python pipeline/analyze_customer_files.py
"""

from __future__ import annotations

import csv
import json
import re
import sys
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

sys.path.insert(0, str(Path(__file__).resolve().parent))
import tax_rules_45x as tx  # noqa: E402

REPO = Path(__file__).resolve().parents[1]
DATA = REPO / "data" / "volterra_synthetic" / "volterra_synthetic_data"
PUBLIC_PREFIX = "demo-files/volterra"  # where scripts/sync-demo-files.mjs publishes the customer files
TAX_PREFIX = "demo-files/45x"
OUT_JSON = REPO / "lib" / "analysis" / "dashboard_analysis.json"
OUT_MD = REPO / "pipeline" / "ANALYSIS_SUMMARY.md"
TAX_YEAR = 2026


# ------------------------------------------------------------------------------------------
# Guarded file access: the pipeline may only read customer-facing files.
# ------------------------------------------------------------------------------------------

def _path(rel: str) -> Path:
    p = DATA / rel
    if "INTERNAL_TEST_ONLY" in p.parts:
        raise PermissionError("The pipeline must not read internal test-oracle files")
    return p


def read_csv(rel: str) -> list[dict]:
    with _path(rel).open(newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def read_xlsx(rel: str, sheet: str | None = None) -> list[dict]:
    wb = load_workbook(_path(rel), read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    # keep the spreadsheet row number (header = row 1) for provenance
    return [dict(zip(rows[0], r), _row=i + 2) for i, r in enumerate(rows[1:])]


def pdf_pages(rel: str) -> list[str]:
    return [" ".join((p.extract_text() or "").split()) for p in PdfReader(str(_path(rel))).pages]


def as_date(v) -> date | None:
    if v in (None, ""):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v)
    if re.fullmatch(r"\d{2}/\d{2}/\d{4}", s):
        return datetime.strptime(s, "%m/%d/%Y").date()
    return date.fromisoformat(s[:10])


# ------------------------------------------------------------------------------------------
# Source files and evidence registry
# ------------------------------------------------------------------------------------------

SOURCE_LABELS = {
    "COMPANY": "Company records",
    "PLM": "Product engineering system (PLM) export",
    "MES": "Factory production system (MES) export",
    "ERP": "Business system (ERP) export",
    "QUALITY": "Engineering / quality test report",
    "SUPPLIER_EVIDENCE": "Supplier document",
}
USED_FOR = {
    "README.md": "Overview of the data package",
    "company_and_facility.xlsx": "Claimant, facility address and plant identifiers",
    "product_master.xlsx": "Product definitions, revisions and specifications",
    "bom.xlsx": "What goes into each cell, by revision",
    "item_master.csv": "Links factory item codes to product numbers",
    "production_orders.csv": "Cells produced, dates and plant",
    "material_consumption.csv": "Which supplier material lots went into each production lot",
    "supplier_master.xlsx": "Supplier names, countries and materials supplied",
    "material_purchase_prices.xlsx": "Material costs per unit",
    "sales.csv": "Cells sold, customers and sale dates",
}

FILES: list[dict] = []
FILE_BY_REL: dict[str, dict] = {}
EVIDENCE: list[dict] = []


def register_files() -> None:
    order = {"COMPANY": 0, "PLM": 1, "MES": 2, "ERP": 3, "QUALITY": 4, "SUPPLIER_EVIDENCE": 5}
    for p in sorted(DATA.rglob("*"), key=lambda x: (order.get(x.relative_to(DATA).parts[0], 9), str(x))):
        if not p.is_file() or "INTERNAL_TEST_ONLY" in p.parts or p.name.startswith("."):
            continue
        rel = p.relative_to(DATA).as_posix()
        folder = rel.split("/")[0] if "/" in rel else ""
        used = USED_FOR.get(p.name)
        if used is None and folder == "QUALITY":
            m = re.match(r"(VX[A-Z0-9]+)_Rev([A-Z])", p.name)
            used = f"Technical specifications and capacity test for {m.group(1).replace('VX', 'VX-')} revision {m.group(2)}" if m else "Product test results"
        if used is None and folder == "SUPPLIER_EVIDENCE":
            used = "Supplier ownership, origin and coverage dates"
        f = {"id": f"F-{len(FILES) + 1:02d}", "path": rel, "name": p.name, "source": SOURCE_LABELS.get(folder, "Data package"),
             "used_for": used, "url": f"{PUBLIC_PREFIX}/{rel}", "used_in": []}
        FILES.append(f)
        FILE_BY_REL[rel] = f


def ev(rel: str, location: str, fact: str, excerpt: str | None = None) -> str:
    """Record a piece of customer evidence and return its id."""
    for e in EVIDENCE:  # de-duplicate identical evidence
        if e["file_id"] == FILE_BY_REL[rel]["id"] and e["location"] == location and e["fact"] == fact:
            return e["id"]
    eid = f"E-{len(EVIDENCE) + 1:03d}"
    EVIDENCE.append({"id": eid, "file_id": FILE_BY_REL[rel]["id"], "location": location, "fact": fact, "excerpt": excerpt})
    return eid


def plain_material(desc: str) -> str:
    """Turn an ERP material description into a plain-English name."""
    d = desc.upper()
    if "CATHODE" in d:
        return "Cathode active material (" + ("NMC811" if "NMC" in d else "LFP" if "LFP" in d else "other") + ")"
    if "ANODE" in d:
        return "Anode active material (synthetic graphite)" if "GRAPHITE" in d else "Anode active material"
    if "SEPARATOR" in d:
        um = re.search(r"(\d+)UM", d)
        return f"Separator ({um.group(1)} µm{', ceramic-coated' if 'CERAMIC' in d else ''})" if um else "Separator"
    if "ELECTROLYTE" in d:
        return "Electrolyte"
    if "CAN" in d and "CAP" in d:
        return "Cell can and cap"
    if "CASE" in d:
        return "Prismatic aluminum case"
    return desc.capitalize()


def fmt_day(iso: str) -> str:
    dt = datetime.strptime(iso, "%Y-%m-%d")
    return f"{dt:%B} {dt.day}, {dt.year}"


def money(x: float) -> str:
    return f"${x:,.0f}"


# ------------------------------------------------------------------------------------------
# Analysis
# ------------------------------------------------------------------------------------------

def analyze() -> dict:
    register_files()

    # ---------------- company and facility
    le_rows = read_xlsx("COMPANY/company_and_facility.xlsx", "LegalEntity")
    le = {r["FIELD"]: (r["VALUE"], r["_row"]) for r in le_rows}
    fac_rows = read_xlsx("COMPANY/company_and_facility.xlsx", "Facilities")
    fac = fac_rows[0]
    cf = "COMPANY/company_and_facility.xlsx"
    company = {
        "name": le["Legal name"][0], "entity_type": le["Entity type"][0], "state": le["State of incorporation"][0],
        "ein": le["EIN (SYNTHETIC)"][0], "tax_year": TAX_YEAR,
        "parent": le["Parent company"][0], "affiliates": le["Subsidiaries / affiliates"][0],
        "evidence": {
            "name": ev(cf, f"Sheet LegalEntity, row {le['Legal name'][1]}", f"Legal name: {le['Legal name'][0]}", le["Legal name"][0]),
            "entity": ev(cf, f"Sheet LegalEntity, rows {le['Entity type'][1]}-{le['State of incorporation'][1]}",
                         f"{le['Entity type'][0]} incorporated in {le['State of incorporation'][0]}"),
            "parent": ev(cf, f"Sheet LegalEntity, row {le['Parent company'][1]}", f"Parent company: {le['Parent company'][0]}", str(le["Parent company"][0])),
            "affiliates": ev(cf, f"Sheet LegalEntity, row {le['Subsidiaries / affiliates'][1]}",
                             f"Subsidiaries / affiliates: {le['Subsidiaries / affiliates'][0]}", str(le["Subsidiaries / affiliates"][0])),
        },
    }
    mentions_48c = any("48C" in str(v[0]) or "48C" in str(k) for k, v in le.items()) or any("48C" in str(v) for v in fac.values())
    facility = {
        "name": fac["FACILITY_NAME"], "address": f"{fac['STREET']}, {fac['CITY']}, {fac['STATE']} {fac['ZIP']}", "country": fac["COUNTRY"],
        "owner": fac["OWNER"], "operator": fac["OPERATOR"], "erp_plant": fac["ERP_PLANT"], "mes_site": fac["MES_SITE"],
        "evidence": ev(cf, f"Sheet Facilities, row {fac['_row']}",
                       f"{fac['FACILITY_NAME']}, {fac['CITY']}, {fac['STATE']}, {fac['COUNTRY']}; owned and operated by {fac['OPERATOR']}; "
                       f"ERP plant {fac['ERP_PLANT']}, MES site {fac['MES_SITE']}"),
    }

    # ---------------- products, revisions, identifiers
    pm = read_xlsx("PLM/product_master.xlsx")
    items = read_csv("MES/item_master.csv")
    pmf = "PLM/product_master.xlsx"
    mes_to_plm = {r["item_code"]: r["plm_part_ref"] for r in items}
    products: dict[str, dict] = {}
    for r in pm:
        pid = r["PART_NUMBER"]
        p = products.setdefault(pid, {"id": pid, "name": r["PART_NAME"], "chemistry": r["CHEMISTRY"], "form_factor": r["FORM_FACTOR"],
                                      "erp_material": r["ERP_MATERIAL_NO"], "revisions": [], "plm_rows": {}})
        p["revisions"].append({"rev": r["REVISION"], "from": as_date(r["EFFECTIVE_FROM"]).isoformat(),
                               "to": as_date(r["EFFECTIVE_TO"]).isoformat() if r["EFFECTIVE_TO"] else None,
                               "change": r["CHANGE_DESCRIPTION"], "change_notice": r["CHANGE_NOTICE"],
                               "evidence": ev(pmf, f"Sheet Products, row {r['_row']}",
                                              f"{pid} revision {r['REVISION']} effective {as_date(r['EFFECTIVE_FROM'])}"
                                              f"{' to ' + str(as_date(r['EFFECTIVE_TO'])) if r['EFFECTIVE_TO'] else ' onward'} ({r['CHANGE_NOTICE']})")})
        p["plm_rows"][r["REVISION"]] = r
        p["mes_item"] = next((m for m, q in mes_to_plm.items() if q == pid), None)
    for p in products.values():
        p["identity_evidence"] = [
            ev(pmf, f"Sheet Products, column ERP_MATERIAL_NO (part {p['id']})", f"{p['id']} = ERP material {p['erp_material']}", p["erp_material"]),
            ev("MES/item_master.csv", f"Record item_code={p['mes_item']}", f"MES item {p['mes_item']} refers to PLM part {p['id']}", p["id"]),
        ]

    def rev_on(pid: str, day: date) -> str:
        for rv in products[pid]["revisions"]:
            if date.fromisoformat(rv["from"]) <= day and (rv["to"] is None or day <= date.fromisoformat(rv["to"])):
                return rv["rev"]
        raise ValueError((pid, day))

    # ---------------- qualification test reports
    reports: dict[tuple[str, str], dict] = {}
    for f in FILES:
        if not f["path"].startswith("QUALITY/"):
            continue
        rel = f["path"]
        pg = pdf_pages(rel)
        m = re.search(r"Product: (\S+) Revision (\S+)", pg[0])
        pid, rev = m.group(1), m.group(2)
        status = re.search(r"Document status: (.+?) Product:", pg[0]).group(1)
        rep = {
            "file": rel, "product": pid, "rev": rev, "approved": status.lower() == "approved", "status": status,
            "standard": "IEC 61960-3" if "IEC 61960-3" in pg[0] else None,
            "voltage_v": float(re.search(r"Nominal voltage: ([\d.]+) V", pg[1]).group(1)),
            "volume_ml": float(re.search(r"External volume: ([\d.]+) mL", pg[1]).group(1)),
            "rated_ah": float(re.search(r"Declared rated capacity: ([\d.]+) Ah", pg[2]).group(1)),
            "rated_wh": float(re.search(r"Rated energy \(nominal voltage x rated capacity\): ([\d.]+) Wh", pg[2]).group(1)),
            "density_wh_l": float(re.search(r"Volumetric energy density \(rated energy / external volume\): ([\d.]+) Wh/L", pg[2]).group(1)),
            "max_power_w": float(re.search(r"Maximum continuous discharge power at nominal voltage: ([\d.]+) W", pg[3]).group(1)),
            "cpr": float(re.search(r"Capacity-to-power ratio \(rated energy / max continuous power\): ([\d.]+) : 1", pg[3]).group(1)),
            "reviewer_pending": "(pending)" in pg[3],
        }
        rep["ev"] = {
            "scope": ev(rel, "Page 1", f"Report covers {pid} revision {rev}; status: {status}", f"Product: {pid} Revision {rev}"),
            "standard": ev(rel, "Page 1, Test procedures", "Capacity measured under IEC 61960-3", "IEC 61960-3 (rated capacity)"),
            "voltage": ev(rel, "Page 2, Section 1", f"Nominal voltage {rep['voltage_v']} V", f"Nominal voltage: {rep['voltage_v']} V"),
            "volume": ev(rel, "Page 2, Section 1", f"External volume {rep['volume_ml']} mL", f"External volume: {rep['volume_ml']:.2f} mL"),
            "capacity": ev(rel, "Page 3, Section 2", f"Declared rated capacity {rep['rated_ah']} Ah", f"Declared rated capacity: {rep['rated_ah']} Ah"),
            "energy": ev(rel, "Page 3, Section 2", f"Rated energy {rep['rated_wh']} Wh per cell",
                         f"Rated energy (nominal voltage x rated capacity): {rep['rated_wh']:.2f} Wh"),
            "density": ev(rel, "Page 3, Section 2", f"Energy density {rep['density_wh_l']:.0f} Wh/L",
                          f"Volumetric energy density (rated energy / external volume): {rep['density_wh_l']:.0f} Wh/L"),
            "cpr": ev(rel, "Page 4, Section 3", f"Capacity-to-power ratio {rep['cpr']} : 1",
                      f"Capacity-to-power ratio (rated energy / max continuous power): {rep['cpr']:.2f} : 1"),
            "signoff": ev(rel, "Page 4, Section 4", "Reviewer signature pending - report not approved" if rep["reviewer_pending"] else "Report signed and approved",
                          "Reviewer: ________________ (pending)" if rep["reviewer_pending"] else None),
        }
        reports[(pid, rev)] = rep

    # ---------------- BOM, suppliers, supplier documents
    bom = read_xlsx("PLM/bom.xlsx")
    vendors = read_xlsx("ERP/supplier_master.xlsx")
    prices = {r["MATNR"]: r for r in read_xlsx("ERP/material_purchase_prices.xlsx")}
    vendor_by_mat = {}
    for v in vendors:
        for m in v["MATERIALS_SUPPLIED"].split("; "):
            vendor_by_mat[m.split(" ")[0]] = v
    sup_docs: dict[str, list[dict]] = defaultdict(list)
    origin: dict[str, dict] = {}
    for f in FILES:
        if not f["path"].startswith("SUPPLIER_EVIDENCE/"):
            continue
        rel = f["path"]
        pg = pdf_pages(rel)
        nm = re.search(r"Supplier legal name: (.+?) Registered address", pg[0])
        if not nm:
            exp = re.search(r"Exporter: (.+?), ", pg[0])
            coo = re.search(r"Country of origin of goods: (\w+)", pg[0])
            if exp and coo:
                origin[exp.group(1)] = {"country": coo.group(1), "ev": ev(rel, "Page 1", f"Country of origin: {coo.group(1)}",
                                                                          f"Country of origin of goods: {coo.group(1)}")}
            continue
        name = nm.group(1)
        kind = re.search(r"(Supplier \w+) to Volterra", pg[0]).group(1)
        up = re.search(r"Ultimate parent: (.+?) Coverage", pg[0]).group(1)
        ip = re.search(r"Immediate parent: (.+?) Ultimate parent", pg[0]).group(1)
        cov = re.search(r"from (\d{4}-\d{2}-\d{2}) through (\d{4}-\d{2}-\d{2})", pg[0])
        sites = {}
        for v in vendors:
            if v["NAME1"] == name:
                for m in v["MATERIALS_SUPPLIED"].split("; "):
                    mat, desc = m.split(" ", 1)
                    sm = re.search(re.escape(mat) + " " + re.escape(desc) + r" (.+?)(?= MAT-| Ownership)", pg[0])
                    if sm:
                        sites[mat] = sm.group(1)
        sup_docs[name].append({
            "file": rel, "kind": kind, "ultimate_parent": up, "immediate_parent": ip,
            "from": date.fromisoformat(cov.group(1)), "to": date.fromisoformat(cov.group(2)),
            "perjury": "under penalty of perjury" in pg[1], "sites": sites,
            "ev": {"up": ev(rel, "Page 1, Ownership", f"Ultimate parent: {up}", f"Ultimate parent: {up}"),
                   "ip": ev(rel, "Page 1, Ownership", f"Immediate parent: {ip}", f"Immediate parent: {ip}"),
                   "cov": ev(rel, "Page 1, Coverage", f"Covers deliveries {cov.group(1)} to {cov.group(2)}",
                             f"Coverage period: materials delivered to Volterra from {cov.group(1)} through {cov.group(2)}"),
                   "sites": ev(rel, "Page 1, Materials covered", "Manufacturing sites: " + "; ".join(f"{k} at {v}" for k, v in sites.items()))},
        })

    # ---------------- production and consumption
    orders = read_csv("MES/production_orders.csv")
    cons = read_csv("MES/material_consumption.csv")
    lot_qty, lot_cons = {}, defaultdict(list)
    for c in cons:
        lot_qty.setdefault(c["fg_lot_id"], int(c["fg_lot_qty"]))
        lot_cons[c["fg_lot_id"]].append(c)
    lots = {}
    for o in orders:
        pid = mes_to_plm[o["item_code"]]
        start = as_date(o["start_ts"])
        lots[o["lot_id"]] = {"lot": o["lot_id"], "order": o["order_id"], "product": pid, "rev": rev_on(pid, start), "start": start,
                             "end": as_date(o["end_ts"]), "site": o["site"], "reported_good": int(o["qty_good"]),
                             "lot_good": lot_qty[o["lot_id"]], "scrap": int(o["qty_scrap"]),
                             "ev": ev("MES/production_orders.csv", f"Record order_id={o['order_id']}",
                                      f"Order {o['order_id']} (lot {o['lot_id']}): {int(o['qty_good']):,} good cells, "
                                      f"{start} to {as_date(o['end_ts'])}, site {o['site']}"),
                             "lot_ev": ev("MES/material_consumption.csv", f"Records fg_lot_id={o['lot_id']}",
                                          f"Lot {o['lot_id']} finished quantity {lot_qty[o['lot_id']]:,} cells")}

    # Lot-level evidence gaps (computed from the customer files only)
    lot_flags: dict[str, set] = defaultdict(set)
    flag_detail: dict[str, dict] = defaultdict(dict)
    for lot, L in lots.items():
        rep = reports.get((L["product"], L["rev"]))
        if rep is None:
            lot_flags[lot].add("NO_TEST")
        elif not rep["approved"]:
            lot_flags[lot].add("DRAFT_TEST")
        if L["reported_good"] != L["lot_good"]:
            lot_flags[lot].add("QTY")
        for c in lot_cons[lot]:
            v = vendor_by_mat[c["component_material"]]
            docs = sup_docs.get(v["NAME1"], [])
            rec = as_date(c["component_lot_received"])
            if not any(dd["from"] <= rec <= dd["to"] for dd in docs):
                lot_flags[lot].add(f"UNCOVERED:{v['NAME1']}")
                flag_detail[f"UNCOVERED:{v['NAME1']}"].setdefault(lot, []).append((c["component_lot"], rec))
            if docs and all(dd["ultimate_parent"].lower().startswith("not disclosed") for dd in docs):
                lot_flags[lot].add(f"OWNERSHIP:{v['NAME1']}")

    # ---------------- capacity basis per product (rule R-CAPACITY-MEASURE / R-CELL-RATE)
    capacity = {}
    for pid, p in products.items():
        approved = [r for (rp, _), r in reports.items() if rp == pid and r["approved"]]
        drafts = [r for (rp, _), r in reports.items() if rp == pid and not r["approved"]]
        latest_plm = p["plm_rows"][p["revisions"][-1]["rev"]]
        plm_wh = round(latest_plm["NOMINAL_VOLTAGE_V"] * latest_plm["RATED_CAPACITY_AH"], 2)
        if approved:
            r = approved[0]
            capacity[pid] = {"basis": "approved_test", "wh": r["rated_wh"], "upper_wh": r["rated_wh"], "report": r,
                             "plm_wh": plm_wh, "conflict": abs(plm_wh - r["rated_wh"]) > 1e-6}
        elif drafts:
            r = drafts[0]
            capacity[pid] = {"basis": "draft_test", "wh": min(r["rated_wh"], plm_wh), "upper_wh": max(r["rated_wh"], plm_wh), "report": r,
                             "plm_wh": plm_wh, "conflict": abs(plm_wh - r["rated_wh"]) > 1e-6}
        else:
            capacity[pid] = {"basis": "plm_only", "wh": plm_wh, "upper_wh": plm_wh, "report": None, "plm_wh": plm_wh, "conflict": False}
        capacity[pid]["plm_ev"] = ev(pmf, f"Sheet Products, row {latest_plm['_row']}, columns NOMINAL_VOLTAGE_V / RATED_CAPACITY_AH",
                                     f"PLM rated capacity {latest_plm['RATED_CAPACITY_AH']} Ah at {latest_plm['NOMINAL_VOLTAGE_V']} V "
                                     f"({plm_wh} Wh)", f"{latest_plm['RATED_CAPACITY_AH']}")

    # ---------------- sales
    sales = read_csv("ERP/sales.csv")
    erp_to_plm = {p["erp_material"]: pid for pid, p in products.items()}
    customers = sorted({s["NAME1"] for s in sales})
    lines_2026, lines_later, unmatched = [], [], []
    for s in sales:
        s["_date"], s["_qty"] = as_date(s["FKDAT"]), int(s["FKIMG"])
        s["_product"] = erp_to_plm.get(s["MATNR"])
        s["_ev"] = ev("ERP/sales.csv", f"Invoice {s['VBELN']} line {s['POSNR']}",
                      f"{s['FKDAT']}: {s['_qty']:,} x {s['MATNR']} to {s['NAME1']} (batch {s['CHARG'] or 'none'})")
        if s["_product"] is None:
            unmatched.append(s)
        elif s["_date"].year == TAX_YEAR:
            lines_2026.append(s)
        else:
            lines_later.append(s)

    # ---------------- credit per product (R-CELL-RATE, R-AGGREGATE, R-PHASEOUT)
    phase = tx.PHASE_OUT[TAX_YEAR]
    credit = {}
    for pid in products:
        rows = [s for s in lines_2026 if s["_product"] == pid]
        units = sum(s["_qty"] for s in rows)
        cap = capacity[pid]
        kwh = cap["wh"] / 1000
        amount = round(units * kwh * tx.CELL_RATE_PER_KWH * phase, 2)
        upper = round(units * cap["upper_wh"] / 1000 * tx.CELL_RATE_PER_KWH * phase, 2)
        line_credit = {s["VBELN"] + s["POSNR"]: s["_qty"] * kwh * tx.CELL_RATE_PER_KWH * phase for s in rows}
        credit[pid] = {"units": units, "kwh_per_cell": kwh, "total_kwh": round(units * kwh, 3), "amount": amount, "upper": upper,
                       "rows": rows, "line_credit": line_credit, "invoices": len({s["VBELN"] for s in rows}),
                       "customers": sorted({s["NAME1"] for s in rows})}

    def credit_on_lots(pred) -> float:
        return round(sum(credit[s["_product"]]["line_credit"][s["VBELN"] + s["POSNR"]] for s in lines_2026 if pred(s["CHARG"])), 2)

    total = round(sum(c["amount"] for c in credit.values()), 2)
    total_upper = round(sum(c["upper"] for c in credit.values()), 2)

    # ---------------- issues (derived only from the checks above)
    issues = []

    def add_issue(**kw):
        kw["id"] = kw.get("id") or f"I-{len(issues) + 1:02d}"
        issues.append(kw)
        return kw["id"]

    # material assistance - rule method unavailable + supplier evidence gaps (always reviewed)
    gap_suppliers = sorted({f.split(":", 1)[1] for fl in lot_flags.values() for f in fl if f.startswith(("UNCOVERED:", "OWNERSHIP:"))})
    ma_id = add_issue(
        id="I-MATERIAL-ASSISTANCE", severity="CRITICAL", type="SUPPLY CHAIN / SUPPLIER",
        title="Foreign-entity sourcing test has not been performed",
        description="For 2026, a cell does not qualify if it includes 'material assistance' from a prohibited foreign entity. The documents "
                    "available to Marigold state this restriction but not the method for testing it, and supplier evidence is incomplete "
                    f"for {len(gap_suppliers)} supplier(s).",
        why="If the test fails for a product, none of that product's cells qualify.",
        affected=list(products) + gap_suppliers,
        credit_affected=total,
        next_action="Have a tax advisor apply current IRS guidance on material assistance using the bill of materials, material costs and supplier "
                    "documents already provided, after the supplier gaps below are closed.",
        evidence_refs=[ev("PLM/bom.xlsx", "Sheet BOM, all purchased lines", "Bill of materials by revision"),
                       ev("ERP/material_purchase_prices.xlsx", "Sheet PriceList, all rows", "2026 material purchase prices")],
        rule_refs=["R-MATERIAL-ASSISTANCE"])

    for (pid, rev) in sorted({(L["product"], L["rev"]) for L in lots.values()}):
        if (pid, rev) in reports:
            continue
        affected_lots = [l for l, L in lots.items() if L["product"] == pid and L["rev"] == rev]
        covered = [r for (rp, rr), r in reports.items() if rp == pid]
        rv = next(r for r in products[pid]["revisions"] if r["rev"] == rev)
        add_issue(
            severity="HIGH", type="PRODUCT / TECHNICAL", title=f"No capacity test for {pid} revision {rev}",
            description=f"Testing on file covers {', '.join('revision ' + r['rev'] for r in covered)} only. Revision {rev} became effective "
                        f"{rv['from']} ({rv['change']}). No evidence shows whether the earlier test still applies.",
            why="Capacity must be measured under a recognized standard. It sets both whether the cell qualifies and how much credit it earns.",
            affected=[f"{pid} revision {rev}"] + affected_lots,
            credit_affected=credit_on_lots(lambda b, al=affected_lots: b in al),
            next_action=f"Obtain a capacity test report for {pid} revision {rev}, or an engineering assessment that revision {rev} does not change "
                        "the tested capacity.",
            evidence_refs=[r["ev"]["scope"] for r in covered] + [rv["evidence"]],
            rule_refs=["R-CAPACITY-MEASURE", "R-CELL-DEF"])

    for pid, cap in capacity.items():
        if cap["basis"] == "draft_test" or cap["conflict"]:
            r = cap["report"]
            add_issue(
                severity="HIGH", type="PRODUCT / TECHNICAL", title=f"{pid} capacity is not settled",
                description=f"The product record lists {cap['plm_wh'] / products[pid]['plm_rows'][r['rev']]['NOMINAL_VOLTAGE_V']:.1f} Ah "
                            f"({cap['plm_wh']:.1f} Wh), while the test report shows {r['rated_ah']} Ah ({r['rated_wh']:.2f} Wh). "
                            + ("The test report is an unapproved draft." if not r["approved"] else ""),
                why=f"Capacity drives the credit. The two values give {money(credit[pid]['amount'])} and {money(credit[pid]['upper'])}; "
                    "Marigold uses the lower value until this is resolved.",
                affected=[pid] + [l for l, L in lots.items() if L["product"] == pid],
                credit_affected=credit[pid]["amount"],
                next_action="Have engineering confirm the rated capacity and approve (sign) the test report, then update the product record if needed.",
                evidence_refs=[cap["plm_ev"], r["ev"]["capacity"], r["ev"]["energy"], r["ev"]["signoff"]],
                rule_refs=["R-CAPACITY-MEASURE", "R-CELL-RATE"])

    for flag in sorted({f for fl in lot_flags.values() for f in fl if f.startswith("OWNERSHIP:")}):
        name = flag.split(":", 1)[1]
        docs = sup_docs[name]
        v = next(v for v in vendors if v["NAME1"] == name)
        affected_lots = sorted(l for l, fl in lot_flags.items() if flag in fl)
        add_issue(
            severity="HIGH", type="SUPPLY CHAIN / SUPPLIER", title=f"{name}: ultimate parent unresolved",
            description=f"{name} identifies its immediate parent ({docs[0]['immediate_parent']}) but states its ultimate parent is "
                        f"'{docs[0]['ultimate_parent']}'. No other customer file identifies the ultimate owner.",
            why="Ownership is needed to judge whether this supplier is a restricted foreign entity for the sourcing test.",
            affected=[name, v["MATERIALS_SUPPLIED"]] + affected_lots,
            credit_affected=credit_on_lots(lambda b, al=affected_lots: b in al),
            next_action=f"Request a full ownership disclosure from {name}, identifying every owner above the immediate parent.",
            evidence_refs=[d["ev"]["ip"] for d in docs] + [d["ev"]["up"] for d in docs]
                          + [ev("ERP/supplier_master.xlsx", f"Sheet Vendors, row {v['_row']}, PARENT_COMPANY_NOTE", f"Procurement note: {v['PARENT_COMPANY_NOTE']}")],
            rule_refs=["R-MATERIAL-ASSISTANCE"])

    for flag in sorted({f for fl in lot_flags.values() for f in fl if f.startswith("UNCOVERED:")}):
        name = flag.split(":", 1)[1]
        docs = sup_docs.get(name, [])
        detail = flag_detail[flag]
        received = sorted({rec for lst in detail.values() for _, rec in lst})
        add_issue(
            severity="HIGH", type="SUPPLY CHAIN / SUPPLIER",
            title=f"{name}: supplier attestation expired" if docs else f"{name}: no supplier attestation",
            description=(f"The attestation on file covers deliveries through {max(d['to'] for d in docs)}. "
                         f"{sum(len(v) for v in detail.values())} material deliveries received {received[0]} to {received[-1]} are not covered, "
                         f"and they were used in {len(detail)} production lots." if docs else "No supplier attestation was provided."),
            why="Without a current attestation, the supplier's representations do not cover material used in these lots.",
            affected=[name] + sorted(detail),
            credit_affected=credit_on_lots(lambda b, al=list(detail): b in al),
            next_action=f"Obtain a renewed attestation from {name} covering deliveries from {max(d['to'] for d in docs) if docs else 'January 1, 2026'} onward.",
            evidence_refs=[d["ev"]["cov"] for d in docs] + [ev("MES/material_consumption.csv",
                                                                f"Records with component_lot_received after {max(d['to'] for d in docs) if docs else ''}",
                                                                f"{name} lots received {received[0]} to {received[-1]} used in lots {', '.join(sorted(detail))}")],
            rule_refs=["R-MATERIAL-ASSISTANCE"])

    if not mentions_48c:
        add_issue(
            severity="MEDIUM", type="CORPORATE / CLAIMANT", title="§48C status of the Holland plant not documented",
            description="Nothing in the company records says whether any equipment at the plant received a §48C investment credit after August 16, 2022.",
            why="Cells made with such equipment cannot also earn 45X credit. The answer is usually a quick confirmation.",
            affected=[facility["name"]], credit_affected=total,
            next_action="Confirm in writing that no property at the Holland plant received a §48C credit after August 16, 2022.",
            evidence_refs=[facility["evidence"], company["evidence"]["entity"]], rule_refs=["R-48C"])

    for s in unmatched:
        add_issue(
            severity="MEDIUM", type="SALES / TRANSACTIONS", title=f"Unidentified product on invoice {s['VBELN']}",
            description=f"{s['_qty']:,} units of '{s['ARKTX']}' (material {s['MATNR']}) were sold to {s['NAME1']} on {s['FKDAT']}. That material "
                        "number does not match either product, and no production lot is recorded.",
            why="The sale cannot be tied to a qualifying product or lot, so it is left out of the estimate.",
            affected=[f"Invoice {s['VBELN']}", s["NAME1"]], credit_affected=0.0,
            next_action="Identify what was sold and which lot it came from. If they were eligible cells, add the lot reference so they can be evaluated.",
            evidence_refs=[s["_ev"]], rule_refs=["R-SCOPE"])

    for lot, L in lots.items():
        if "QTY" in lot_flags[lot]:
            add_issue(
                severity="LOW", type="PRODUCTION / RECONCILIATION", title=f"Production count mismatch on lot {lot}",
                description=f"The production order reports {L['reported_good']:,} good cells, but the lot record shows {L['lot_good']:,} "
                            f"(a difference of {L['reported_good'] - L['lot_good']}). Material usage matches the lower figure.",
                why="Production counts support how many cells were available to sell. Sales do not exceed the lower figure, so the estimate is unaffected.",
                affected=[L["order"], lot], credit_affected=0.0,
                next_action="Ask the plant to correct the production order or explain the difference.",
                evidence_refs=[L["ev"], L["lot_ev"]], rule_refs=["R-SCOPE"])

    sev_rank = {"CRITICAL": 0, "HIGH": 1, "MEDIUM": 2, "LOW": 3}
    issues.sort(key=lambda i: (sev_rank[i["severity"]], -i["credit_affected"]))
    issues_by_lot = defaultdict(list)
    for i in issues:
        for a in i["affected"]:
            if a in lots:
                issues_by_lot[a].append(i["id"])

    # ---------------- eligibility requirements
    def status_of(ok: bool, review: bool = False) -> str:
        return "satisfied" if ok and not review else ("review" if ok or review else "not_satisfied")

    cell_issues = [i["id"] for i in issues if i["type"] == "PRODUCT / TECHNICAL"]
    def_lines = []
    for pid in products:
        rs = [r for (rp, _), r in reports.items() if rp == pid]
        for r in rs:
            met = r["rated_wh"] >= tx.MIN_ENERGY_WH and r["density_wh_l"] >= tx.MIN_ENERGY_DENSITY_WH_PER_L and r["cpr"] <= tx.MAX_CAPACITY_TO_POWER
            def_lines.append({"product": pid, "rev": r["rev"], "met": met, "approved": r["approved"],
                              "text": f"{pid} rev {r['rev']}: {r['rated_wh']:.1f} Wh, {r['density_wh_l']:.0f} Wh/L, capacity-to-power {r['cpr']:.2f}:1"
                                      f"{'' if r['approved'] else ' (draft report)'}"})
    all_met = all(x["met"] for x in def_lines)
    make_items = [b for b in bom if b["MAKE_BUY"] == "Make"]
    sites = {L["site"] for L in lots.values()}
    requirements = [
        {"id": "ELIG-CELL", "title": "Products are battery cells that meet the technical definition",
         "status": "not_satisfied" if not all_met else ("review" if cell_issues else "satisfied"),
         "explanation": "Each cell must store at least 12 Wh, hold at least 100 Wh per liter, stay within a 100:1 capacity-to-power ratio, and "
                        "have its capacity measured under a recognized standard.",
         "fact": "Every tested product clears the thresholds by a wide margin: " + "; ".join(x["text"] for x in def_lines) +
                 ". " + ("Some test evidence is missing or unapproved." if cell_issues else ""),
         "affected_products": list(products),
         "evidence_refs": [e for r in reports.values() for e in (r["ev"]["energy"], r["ev"]["density"], r["ev"]["cpr"], r["ev"]["standard"])],
         "tax_rule_refs": ["R-CELL-DEF", "R-CPR", "R-CAPACITY-MEASURE"], "issue_refs": cell_issues,
         "next_action": "Approve the VX-LFP50 test report and obtain test evidence for VX-2170 revision C." if cell_issues else None},
        {"id": "ELIG-PRODUCED", "title": "Volterra makes the cells itself",
         "status": "satisfied" if make_items and sites == {facility["mes_site"]} and facility["operator"] == company["name"] else "review",
         "explanation": "The claimant must do the manufacturing that turns materials into a finished cell, not just minor assembly.",
         "fact": f"Volterra makes its own cathode and anode electrodes from purchased materials and assembles the cells itself "
                 f"({len(lots)} production orders in {TAX_YEAR}, all at its own plant).",
         "affected_products": list(products),
         "evidence_refs": [ev("PLM/bom.xlsx", "Sheet BOM, rows with MAKE_BUY = Make", "Electrode subassemblies are made in-house"),
                           facility["evidence"]] + [L["ev"] for L in list(lots.values())[:2]],
         "tax_rule_refs": ["R-PRODUCED-BY", "R-TRADE"], "issue_refs": [], "next_action": None},
        {"id": "ELIG-US", "title": "Cells are made in the United States",
         "status": "satisfied" if facility["country"] == "United States" and sites == {facility["mes_site"]} else "review",
         "explanation": "Cells must be produced in the U.S. Using imported materials is allowed.",
         "fact": f"All {len(lots)} production orders ran at {facility['name']}, {facility['address']}.",
         "affected_products": list(products), "evidence_refs": [facility["evidence"]] + [L["ev"] for L in list(lots.values())[:2]],
         "tax_rule_refs": ["R-US"], "issue_refs": [], "next_action": None},
        {"id": "ELIG-SALES", "title": "Cells were sold in 2026 to unrelated customers",
         "status": "satisfied" if lines_2026 and str(company["affiliates"]).lower() == "none" else "review",
         "explanation": "Only cells sold during the tax year count, and buyers must not be under common control with Volterra.",
         "fact": f"{sum(len({s['VBELN'] for s in lines_2026 if s['_product'] == p}) for p in products)} invoices in 2026 to "
                 f"{len({s['NAME1'] for s in lines_2026})} customers. Company records list no parent or affiliates. "
                 f"{len({s['VBELN'] for s in lines_later})} January 2027 invoices are outside the tax year and excluded.",
         "affected_products": list(products),
         "evidence_refs": [company["evidence"]["parent"], company["evidence"]["affiliates"]] + [s["_ev"] for s in lines_2026[:2]],
         "tax_rule_refs": ["R-SCOPE", "R-UNRELATED"], "issue_refs": [i["id"] for i in issues if i["type"] == "SALES / TRANSACTIONS"],
         "next_action": None},
        {"id": "ELIG-PFE-TAXPAYER", "title": "Volterra is not a restricted foreign entity",
         "status": "satisfied" if str(company["parent"]).lower() == "none" else "review",
         "explanation": "The company claiming the credit cannot itself be a foreign entity of concern or controlled by one.",
         "fact": f"{company['name']} is a {company['state']} {company['entity_type']} with no parent company listed.",
         "affected_products": list(products), "evidence_refs": [company["evidence"]["entity"], company["evidence"]["parent"]],
         "tax_rule_refs": ["R-PFE-TAXPAYER"], "issue_refs": [], "next_action": None},
        {"id": "ELIG-MATERIAL-ASSISTANCE", "title": "Materials do not come from restricted foreign sources",
         "status": "review",
         "explanation": "For 2026, a cell does not qualify if too much of its material is traced to prohibited foreign entities.",
         "fact": "This test has not been performed. The method is not in the documents available, and ownership or attestation evidence "
                 f"is incomplete for: {', '.join(gap_suppliers)}.",
         "affected_products": list(products), "evidence_refs": issues[0]["evidence_refs"],
         "tax_rule_refs": ["R-MATERIAL-ASSISTANCE"], "issue_refs": [i["id"] for i in issues if i["type"] == "SUPPLY CHAIN / SUPPLIER"],
         "next_action": "Close the supplier evidence gaps, then have a tax advisor perform the sourcing test."},
        {"id": "ELIG-48C", "title": "No overlapping §48C investment credit",
         "status": "satisfied" if mentions_48c else "review",
         "explanation": "Equipment that already earned a §48C investment credit cannot also be used to earn 45X credit.",
         "fact": "The company records do not address §48C." if not mentions_48c else "Company records address §48C.",
         "affected_products": list(products), "evidence_refs": [facility["evidence"]], "tax_rule_refs": ["R-48C"],
         "issue_refs": [i["id"] for i in issues if i["type"] == "CORPORATE / CLAIMANT"],
         "next_action": "Confirm no §48C credit was taken for the Holland plant." if not mentions_48c else None},
    ]
    statuses = {r["status"] for r in requirements}
    overall = "not_eligible" if "not_satisfied" in statuses else ("review" if "review" in statuses else "eligible")
    eligibility = {
        "status": overall,
        "headline": {"eligible": "Volterra appears eligible for the Section 45X battery-cell credit",
                     "review": "Volterra appears to qualify, but eligibility cannot yet be fully established",
                     "not_eligible": "Volterra does not currently meet every requirement"}[overall],
        "summary": f"{sum(r['status'] == 'satisfied' for r in requirements)} of {len(requirements)} requirements are satisfied by the files "
                   f"provided. {sum(r['status'] == 'review' for r in requirements)} need more information or review.",
        "requirements": requirements,
    }

    # ---------------- per-product output
    out_products = []
    for pid, p in products.items():
        cap = capacity[pid]
        rep = cap["report"]
        plots = [L for L in lots.values() if L["product"] == pid]
        c = credit[pid]
        rows_by_rev = {}
        for rv in p["revisions"]:
            mats = []
            for b in [b for b in bom if b["TOP_ASSEMBLY"] == pid and b["TOP_REV"] == rv["rev"] and b["MAKE_BUY"] == "Purchased"]:
                v = vendor_by_mat[b["ERP_MATERIAL_NO"]]
                docs = sup_docs.get(v["NAME1"], [])
                site = next((d["sites"].get(b["ERP_MATERIAL_NO"]) for d in docs if d["sites"].get(b["ERP_MATERIAL_NO"])), None)
                ids = [i["id"] for i in issues if v["NAME1"] in i["affected"] and i["id"] != ma_id]
                rev_lots = [L["lot"] for L in plots if L["rev"] == rv["rev"]]
                ids = [i for i in ids if any(l in next(x for x in issues if x["id"] == i)["affected"] for l in rev_lots)] or \
                      [i for i in ids if next(x for x in issues if x["id"] == i)["title"].endswith("ultimate parent unresolved")]
                mats.append({
                    "material": plain_material(b["COMPONENT_DESCRIPTION"]), "erp_material": b["ERP_MATERIAL_NO"],
                    "quantity": f"{b['QTY_PER']} {b['UOM'].lower().replace('m2', 'm²')} per cell",
                    "supplier": v["NAME1"], "made_in": site.split(", ")[-1] if site else v["LAND1"],
                    "evidence_status": "issue" if ids else ("supported" if docs else "missing"),
                    "issue_refs": ids,
                    "evidence_refs": [ev("PLM/bom.xlsx", f"Sheet BOM, row {b['_row']}", f"{pid} rev {rv['rev']}: {b['COMPONENT_ITEM']} "
                                                                                         f"{b['QTY_PER']} {b['UOM']} per cell")]
                                     + ([docs[0]["ev"]["sites"], docs[0]["ev"]["cov"]] if docs else []),
                })
            rows_by_rev[rv["rev"]] = mats
        changes = []
        revs = p["revisions"]
        for a, b2 in zip(revs, revs[1:]):
            old = {m["erp_material"]: m for m in rows_by_rev[a["rev"]]}
            new = {m["erp_material"]: m for m in rows_by_rev[b2["rev"]]}
            for k in sorted(new.keys() - old.keys()):
                for k_old in sorted(old.keys() - new.keys()):
                    o, n = old[k_old], new[k]
                    changes.append(f"From {fmt_day(b2['from'])}, revision {b2['rev']} switches the {o['material'].split(' (')[0].lower()} from "
                                   f"{o['supplier']} ({o['material'].split('(')[-1].rstrip(')')}, made in {o['made_in']}) to "
                                   f"{n['supplier']} ({n['material'].split('(')[-1].rstrip(')')}, made in {n['made_in']}).")
        p_issue_ids = [i["id"] for i in issues if pid in i["affected"] or any(l in i["affected"] for l in (L["lot"] for L in plots))]
        specs = []
        if rep:
            specs = [
                {"label": "Nominal voltage", "value": f"{rep['voltage_v']} V", "evidence_refs": [rep["ev"]["voltage"]]},
                {"label": "Rated capacity", "value": f"{rep['rated_ah']} Ah" + (f" (product record: {cap['plm_wh'] / rep['voltage_v']:.1f} Ah)" if cap["conflict"] else ""),
                 "evidence_refs": [rep["ev"]["capacity"]] + ([cap["plm_ev"]] if cap["conflict"] else [])},
                {"label": "Energy per cell", "value": f"{rep['rated_wh']:.2f} Wh", "evidence_refs": [rep["ev"]["energy"]]},
                {"label": "Energy density", "value": f"{rep['density_wh_l']:.0f} Wh/L", "evidence_refs": [rep["ev"]["density"]]},
                {"label": "Capacity-to-power ratio", "value": f"{rep['cpr']:.2f} : 1", "evidence_refs": [rep["ev"]["cpr"]]},
                {"label": "Test report", "value": f"Revision {rep['rev']}, {'approved' if rep['approved'] else 'draft - not approved'}",
                 "evidence_refs": [rep["ev"]["scope"], rep["ev"]["signoff"]]},
            ]
        out_products.append({
            "id": pid, "name": p["name"], "chemistry": p["chemistry"], "form_factor": p["form_factor"],
            "identifiers": {"plm": pid, "erp": p["erp_material"], "mes": p["mes_item"], "evidence_refs": p["identity_evidence"]},
            "revisions": [{"rev": r["rev"], "from": r["from"], "to": r["to"], "change": r["change"], "evidence_refs": [r["evidence"]],
                           "materials": rows_by_rev[r["rev"]]} for r in revs],
            "revision_changes": changes,
            "specs": specs,
            "facility": facility["name"],
            "production": {"cells": sum(L["lot_good"] for L in plots), "orders": len(plots),
                           "first": min(L["start"] for L in plots).isoformat(), "last": max(L["end"] for L in plots).isoformat(),
                           "evidence_refs": [L["ev"] for L in plots]},
            "sales": {"cells": c["units"], "invoices": c["invoices"], "customers": c["customers"],
                      "evidence_refs": [s["_ev"] for s in c["rows"]]},
            "credit": {"amount": c["amount"], "upper": c["upper"]},
            "status": "review" if p_issue_ids else "supported",
            "issue_refs": p_issue_ids,
        })

    # ---------------- credit calculations
    calcs = []
    for pid in products:
        cap, c = capacity[pid], credit[pid]
        rep = cap["report"]
        cap_refs = ([rep["ev"]["energy"], rep["ev"]["standard"]] if rep else []) + ([cap["plm_ev"]] if cap["conflict"] or not rep else [])
        note = None
        if cap["conflict"]:
            note = (f"Product record and test report disagree ({cap['plm_wh']:.2f} Wh vs {rep['rated_wh']:.2f} Wh). The estimate uses the lower "
                    f"value; with the higher value it would be {money(c['upper'])}.")
        elif any(i["title"].startswith(f"No capacity test for {pid}") for i in issues):
            note = f"Later revisions have no capacity test of their own. The tested capacity of revision {rep['rev']} is applied to all cells sold, pending that test."
        calcs.append({
            "product": pid,
            "steps": [
                {"label": "Cells sold to customers in 2026", "value": f"{c['units']:,} cells", "evidence_refs": [s["_ev"] for s in c["rows"]],
                 "tax_rule_refs": ["R-SCOPE"]},
                {"label": "Capacity per cell (from the capacity test)", "value": f"{c['kwh_per_cell'] * 1000:.2f} Wh = {c['kwh_per_cell']:.5f} kWh",
                 "evidence_refs": cap_refs, "tax_rule_refs": ["R-CAPACITY-MEASURE", "R-CELL-RATE"]},
                {"label": "Qualifying capacity", "value": f"{c['units']:,} x {c['kwh_per_cell']:.5f} kWh = {c['total_kwh']:,.2f} kWh",
                 "evidence_refs": [], "tax_rule_refs": ["R-AGGREGATE"]},
                {"label": "Credit rate for battery cells", "value": "$35 per kWh", "evidence_refs": [], "tax_rule_refs": ["R-CELL-RATE"]},
                {"label": "2026 phase-out", "value": "100% (no reduction before 2030)", "evidence_refs": [], "tax_rule_refs": ["R-PHASEOUT"]},
            ],
            "formula": f"{c['total_kwh']:,.2f} kWh x $35 x 100% = {money(c['amount'])}",
            "amount": c["amount"], "upper": c["upper"], "note": note,
            "plain": f"Volterra sold {c['units']:,} {pid} cells in 2026. Each stores {c['kwh_per_cell'] * 1000:.2f} Wh, so together they hold "
                     f"{c['total_kwh']:,.0f} kWh. Section 45X pays $35 per kWh, which gives about {money(c['amount'])}.",
        })

    excluded = []
    for reason, group in (("Sold in January 2027, outside the 2026 tax year", lines_later), ("Product could not be identified", unmatched)):
        for inv in sorted({s["VBELN"] for s in group}):
            rows = [s for s in group if s["VBELN"] == inv]
            excluded.append({"label": f"Invoice {inv} ({rows[0]['FKDAT']})", "reason": reason, "cells": sum(s["_qty"] for s in rows),
                             "evidence_refs": [s["_ev"] for s in rows]})

    # ---------------- link files to the conclusions that use them
    conclusions = ([("eligibility", r["id"], r["title"], r["evidence_refs"]) for r in requirements]
                   + [("credits", f"credit-{c['product']}", f"{c['product']} credit calculation", [e for s in c["steps"] for e in s["evidence_refs"]]) for c in calcs]
                   + [("products", p["id"], f"{p['id']} product record", p["identifiers"]["evidence_refs"] + [e for s in p["specs"] for e in s["evidence_refs"]]
                       + [e for r in p["revisions"] for m in r["materials"] for e in m["evidence_refs"]] + p["production"]["evidence_refs"]) for p in out_products]
                   + [("missing-information", i["id"], i["title"], i["evidence_refs"]) for i in issues])
    ev_file = {e["id"]: e["file_id"] for e in EVIDENCE}
    for page, cid, title, refs in conclusions:
        for fid in sorted({ev_file[e] for e in refs}):
            f = next(x for x in FILES if x["id"] == fid)
            if not any(u["id"] == cid for u in f["used_in"]):
                f["used_in"].append({"page": page, "id": cid, "title": title})

    sev_counts = {s: sum(i["severity"] == s for i in issues) for s in ("CRITICAL", "HIGH", "MEDIUM", "LOW")}
    return {
        "meta": {"title": "Marigold - Section 45X analysis", "generated_from": "customer-provided files only",
                 "pipeline": "pipeline/analyze_customer_files.py (pre-computed; simulates the future Marigold ingestion pipeline)",
                 "data_as_of": "2027-02-12", "disclaimer": "Demo analysis on fictional data. Not tax advice."},
        "company": {k: v for k, v in company.items() if k != "evidence"} | {"evidence_refs": list(company["evidence"].values())},
        "facility": {k: v for k, v in facility.items() if k != "evidence"} | {"evidence_refs": [facility["evidence"]]},
        "tax_year": TAX_YEAR,
        "eligibility": eligibility,
        "credit_summary": {"estimate": total, "upper": total_upper,
                           "by_product": [{"product": pid, "amount": credit[pid]["amount"], "upper": credit[pid]["upper"]} for pid in products],
                           "basis": "Cells sold to unrelated customers in 2026 x tested capacity (kWh) x $35 per kWh. Where capacity is disputed, the lower value is used.",
                           "excluded_sales": excluded},
        "credit_calculations": calcs,
        "products": out_products,
        "issues": issues,
        "issue_counts": sev_counts,
        "source_files": FILES,
        "evidence": EVIDENCE,
        "tax_rules": [r | {"url": f"{TAX_PREFIX}/{r['file']}"} for r in tx.RULES],
    }


def write_summary(a: dict) -> None:
    c = a["credit_summary"]
    lines = [
        "# Customer-data analysis summary (internal)", "",
        "Produced by `pipeline/analyze_customer_files.py` from customer-provided files only. No `INTERNAL_TEST_ONLY` file was read.", "",
        f"- **Claimant:** {a['company']['name']} ({a['company']['entity_type']}, {a['company']['state']}); parent: {a['company']['parent']}; affiliates: {a['company']['affiliates']}",
        f"- **Facility:** {a['facility']['name']}, {a['facility']['address']}, {a['facility']['country']} (ERP {a['facility']['erp_plant']}, MES {a['facility']['mes_site']})",
        f"- **Eligibility:** {a['eligibility']['headline']}. {a['eligibility']['summary']}",
        f"- **Estimated 2026 credit:** ${c['estimate']:,.2f} (up to ${c['upper']:,.2f})", "",
        "| Product | Revisions | Produced (lot qty) | Sold 2026 | Capacity used | Credit | Upper |", "|---|---|---|---|---|---|---|",
    ]
    for p, calc in zip(a["products"], a["credit_calculations"]):
        lines.append(f"| {p['id']} | {', '.join(r['rev'] for r in p['revisions'])} | {p['production']['cells']:,} | {p['sales']['cells']:,} | "
                     f"{calc['steps'][1]['value']} | ${calc['amount']:,.2f} | ${calc['upper']:,.2f} |")
    lines += ["", "## Requirements", ""] + [f"- `{r['status']}` {r['title']}: {r['fact']}" for r in a["eligibility"]["requirements"]]
    lines += ["", "## Issues", ""] + [f"- **{i['severity']}** [{i['type']}] {i['title']} (credit affected ${i['credit_affected']:,.2f})" for i in a["issues"]]
    lines += ["", "## Excluded sales", ""] + [f"- {x['label']}: {x['cells']:,} cells - {x['reason']}" for x in c["excluded_sales"]]
    OUT_MD.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    a = analyze()
    OUT_JSON.parent.mkdir(parents=True, exist_ok=True)
    OUT_JSON.write_text(json.dumps(a, indent=2, default=str) + "\n", encoding="utf-8")
    write_summary(a)
    print(f"Wrote {OUT_JSON.relative_to(REPO)} ({len(a['evidence'])} evidence items, {len(a['issues'])} issues, {len(a['source_files'])} files)")
    print(f"Estimated 2026 credit ${a['credit_summary']['estimate']:,.2f} (up to ${a['credit_summary']['upper']:,.2f}); "
          f"eligibility: {a['eligibility']['status']}; issues: {a['issue_counts']}")


if __name__ == "__main__":
    main()
